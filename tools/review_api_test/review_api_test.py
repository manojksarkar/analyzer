#!/usr/bin/env python3
"""Review & Update, end to end through the REST API.

Onboards a project and generates a version the way the web app does, then exercises the whole
review feature on it: every kind of correction, every route that reads one, second edits, undo, two
saves of one slot at the same moment, the mistakes a client can make, the re-export and the Word
file, and the next version carrying the corrections. Each answer is checked through the API itself.

Only request bodies come from the config (config.example.json). Every id the server makes --
project, job, version, document, slot key -- is read from its answers.

    python tools/review_api_test/review_api_test.py --config my_config.json

README.md in this folder explains the config and every step. Exit code: 0 every check passed,
1 a check failed, 2 the test could not run (config, sign-in, a generation that failed).
"""
from __future__ import annotations

import argparse
import concurrent.futures
import copy
import datetime as dt
import io
import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
REPO_ROOT = os.path.dirname(os.path.dirname(HERE))

#: What every route answers for a slot (API spec §5 `Slot`), and what a save adds to it.
SLOT_FIELDS = ("slotKind", "slotKey", "text", "llmText", "humanText", "isOverridden",
               "isOrphaned", "canUndo", "updatedBy", "updatedAt")
SAVE_FIELDS = ("previousText", "firstEdit", "viewsDerived", "queuedForRegeneration")
#: Saved one slot at a time through R3. The other two kinds have routes of their own (R6, R8).
R3_KINDS = ("description", "inputName", "outputName", "unitDescription", "structDescription")
ALL_KINDS = R3_KINDS + ("behaviourDescription", "nodeLabel")
#: On the document page straight after a save (API spec, "When a correction becomes visible").
#: A struct description reaches the page with the next re-export.
PAGE_AFTER_SAVE = ("description", "inputName", "outputName", "behaviourDescription", "nodeLabel",
                   "unitDescription")
FINISHED = ("complete", "failed", "cancelled")

DEFAULT_OPTIONS = {
    "kinds": list(ALL_KINDS),
    "slots_per_kind": 2,
    "concurrent_saves": True,
    "negative_checks": True,
    "check_page": True,
    "reexport": True,
    "check_word": True,
    "carry_forward": True,
    "report_file": "review_api_test_report.json",
}


class StopTest(Exception):
    """The test cannot go on: a prerequisite -- sign-in, the project, the version -- failed."""


# ---------------------------------------------------------------------------
# small helpers
# ---------------------------------------------------------------------------
def short(value, limit=300) -> str:
    text = value if isinstance(value, str) else json.dumps(value, ensure_ascii=False, default=str)
    text = " ".join(text.split())
    return text if len(text) <= limit else text[:limit] + " ..."


def norm(text) -> str:
    """Whitespace collapsed -- a Word cell or a page payload may break a line anywhere."""
    return " ".join((text or "").split())


def _svg_urls(node) -> list[str]:
    """Every `image_url` in a render payload that names an SVG -- its flowchart pictures."""
    found: list[str] = []
    if isinstance(node, dict):
        url = node.get("image_url")
        if isinstance(url, str) and url.split("?", 1)[0].lower().endswith(".svg"):
            found.append(url)
        for value in node.values():
            found.extend(_svg_urls(value))
    elif isinstance(node, list):
        for value in node:
            found.extend(_svg_urls(value))
    return found


_PLACEHOLDER = re.compile(r"\{(timestamp|projectId|versionId|versionTag)\}")


def fill(obj, values):
    """`obj` with the `{placeholders}` filled from what the server has answered so far.

    Keys that start with `_` are comments (JSON has none of its own) and are dropped.
    """
    if isinstance(obj, str):
        return _PLACEHOLDER.sub(lambda m: str(values.get(m.group(1), m.group(0))), obj)
    if isinstance(obj, list):
        return [fill(x, values) for x in obj]
    if isinstance(obj, dict):
        return {k: fill(v, values) for k, v in obj.items() if not str(k).startswith("_")}
    return obj


def fields(slot) -> dict:
    return {k: (slot or {}).get(k) for k in SLOT_FIELDS}


def diff(want: dict, got: dict) -> list:
    return ["%s: want %r, got %r" % (k, want.get(k), got.get(k))
            for k in SLOT_FIELDS if want.get(k) != got.get(k)]


def body_of(response) -> dict:
    try:
        out = response.json()
    except ValueError:
        return {}
    return out if isinstance(out, dict) else {}


def docx_text(content: bytes) -> str:
    """Every paragraph and table cell of a Word file, nested tables included."""
    import docx                                    # python-docx, in requirements.txt
    document = docx.Document(io.BytesIO(content))
    out = [p.text for p in document.paragraphs]

    def walk(tables):
        for table in tables:
            for row in table.rows:
                for cell in row.cells:
                    out.append(cell.text)
                    walk(cell.tables)
    walk(document.tables)
    return norm("\n".join(out))


# ---------------------------------------------------------------------------
# results
# ---------------------------------------------------------------------------
class Checks:
    """Every check, printed as it runs and kept for the summary and the JSON report."""

    def __init__(self):
        self.items = []
        self.step = ""

    def begin(self, title):
        self.step = title
        print("\n== %s" % title, flush=True)

    def check(self, name, ok, detail=None, *, warn=False) -> bool:
        result = "PASS" if ok else ("WARN" if warn else "FAIL")
        self.items.append({"step": self.step, "check": name, "result": result,
                           "detail": None if detail is None else short(detail, 3000)})
        line = "   %s  %s" % (result, name)
        if not ok and detail not in (None, "", [], {}):
            line += "\n         " + short(detail, 600)
        print(line, flush=True)
        return ok

    def skip(self, name, why):
        self.items.append({"step": self.step, "check": name, "result": "SKIP", "detail": why})
        print("   SKIP  %s -- %s" % (name, why), flush=True)

    def note(self, text):
        print("         %s" % text, flush=True)

    def counts(self) -> dict:
        out = {"PASS": 0, "FAIL": 0, "WARN": 0, "SKIP": 0}
        for item in self.items:
            out[item["result"]] += 1
        return out


# ---------------------------------------------------------------------------
# the API
# ---------------------------------------------------------------------------
class Api:
    """Signs in, sends requests, and signs in again once when a token has expired."""

    def __init__(self, base_url, email, password, *, timeout=300, verify=True, transport=None):
        base = (base_url or "http://localhost:8000").rstrip("/")
        self.base = base if base.endswith("/api/v1") else base + "/api/v1"
        self.email, self.password, self.timeout = email, password, timeout
        self.in_process = transport is not None
        if transport is None:
            import requests
            transport = requests.Session()
        self.http = transport
        self.extra = {} if self.in_process else {"verify": verify}
        self.token = None

    def _send(self, method, path, *, json=None, params=None, auth=True):
        url = path if path.startswith("http") else self.base + path
        headers = {"Authorization": "Bearer " + self.token} if auth and self.token else {}
        # A new connection for every request. A kept-alive one goes stale while a job is
        # followed: uvicorn closes it after 5 idle seconds, and a request sent on it just then
        # fails with "Remote end closed connection" although the server never saw it.
        headers["Connection"] = "close"
        # A read is retried on a network error; a write never is -- it may have landed.
        attempts = 3 if method == "GET" else 1
        for attempt in range(attempts):
            try:
                return self.http.request(method, url, json=json, params=params,
                                         headers=headers, timeout=self.timeout, **self.extra)
            except Exception as exc:               # noqa: BLE001 -- say where, then stop
                if not type(exc).__module__.startswith(("requests", "urllib3", "http")):
                    raise
                if attempt + 1 < attempts:
                    time.sleep(1 + attempt)
                    continue
                raise StopTest("cannot reach %s: %s" % (url, exc)) from None

    def signin(self):
        r = self._send("POST", "/auth/signin",
                       json={"email": self.email, "password": self.password}, auth=False)
        if r.status_code != 200:
            raise StopTest("sign-in as %s failed: %s %s" % (self.email, r.status_code,
                                                            short(r.text)))
        self.token = r.json()["access_token"]

    def call(self, method, path, *, json=None, params=None, auth=True):
        r = self._send(method, path, json=json, params=params, auth=auth)
        if r.status_code == 401 and auth and self.token:
            self.signin()
            r = self._send(method, path, json=json, params=params, auth=auth)
        return r

    def twin(self) -> "Api":
        """A second client with the same sign-in, for two requests at the same moment."""
        other = copy.copy(self)
        if not self.in_process:
            import requests
            other.http = requests.Session()
        return other


# ---------------------------------------------------------------------------
# the test
# ---------------------------------------------------------------------------
class ReviewApiTest:

    def __init__(self, api: Api, cfg: dict, checks: Checks, args):
        self.api, self.cfg, self.c = api, cfg, checks
        self.opt = dict(DEFAULT_OPTIONS, **(cfg.get("test") or {}))
        stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
        #: Written into every corrected text, so a text is found only if THIS run wrote it.
        self.marker = "T" + stamp[-6:]
        self.values = {"timestamp": stamp}
        self.pid = args.project_id or (cfg.get("onboard") or {}).get("project_id") or None
        self.vid = args.version_id or (cfg.get("generate") or {}).get("version_id") or None
        if self.vid and not self.pid:
            raise StopTest("a version id was given without its project id")
        #: The version was generated by this run, so it starts with no corrections.
        self.fresh = not self.vid
        self.database = ((cfg.get("server") or {}).get("database") or "postgresql").lower()
        self.ids = {}
        self.targets = []
        self.catalog = {}
        self.job_body = None
        self.flowchart = None
        self.base = ""

    # ---- paths and small reads ----------------------------------------------------------------
    def P(self, path="") -> str:
        return "/projects/%s%s" % (self.pid, path)

    def V(self, path="", version=None) -> str:
        return "/projects/%s/versions/%s%s" % (self.pid, version or self.vid, path)

    def slots(self, kind, version=None, **params) -> list:
        r = self.api.call("GET", self.V("/slots", version),
                          params={"slot_kind": kind, "limit": 1000, **params})
        if r.status_code != 200:
            raise StopTest("R11 %s: %s %s" % (kind, r.status_code, short(r.text)))
        return r.json().get("slots") or []

    def read(self, t, version=None):
        return self.api.call("GET", self.V("/overrides/slot", version),
                             params={"slot_kind": t["kind"], "slot_key": t["key"]})

    def history(self, t) -> list:
        r = self.api.call("GET", self.V("/overrides/history"),
                          params={"slot_kind": t["kind"], "slot_key": t["key"]})
        return [h.get("humanText") for h in body_of(r).get("history") or []]

    def readiness(self) -> dict:
        return body_of(self.api.call("GET", self.V("/export-readiness")))

    def queued(self, version=None) -> set:
        r = self.api.call("GET", self.V("/regeneration-queue", version))
        return {(p.get("slotKind"), p.get("slotKey")) for p in body_of(r).get("pending") or []}

    def labels(self, flowchart_id, version=None) -> dict:
        r = self.api.call("GET", self.V("/flowcharts/labels", version),
                          params={"flowchart_id": flowchart_id})
        return body_of(r)

    def documents(self, version=None) -> list:
        r = self.api.call("GET", self.P("/documents"),
                          params={"version_id": version or self.vid, "per_page": 100})
        return body_of(r).get("documents") or []

    def wait_job(self, job_id, what):
        gen = self.cfg.get("generate") or {}
        poll = float(gen.get("poll_seconds") or 5)
        limit = float(gen.get("timeout_minutes") or 60) * 60
        started, last = time.time(), None
        while True:
            r = self.api.call("GET", self.P("/jobs/%s" % job_id))
            if r.status_code != 200:
                raise StopTest("GET job %s: %s %s" % (job_id, r.status_code, short(r.text)))
            job = r.json()["job"]
            line = "%s %s: %s, phase %s" % (what, job_id, job.get("status"), job.get("phase"))
            if line != last:
                self.c.note(line + (" -- " + short(job.get("current_activity") or "", 90)
                                    if job.get("current_activity") else ""))
                last = line
            if job.get("status") in FINISHED:
                return job, int(time.time() - started)
            if time.time() - started > limit:
                raise StopTest("%s %s did not finish in %d minutes" % (what, job_id, limit // 60))
            time.sleep(poll)

    # ---- the targets: the slots this run corrects ---------------------------------------------
    def by_kind(self, *kinds) -> list:
        return [t for t in self.targets if t["kind"] in kinds]

    def name(self, t) -> str:
        if t["kind"] == "nodeLabel":
            return "nodeLabel [%s %s]" % (t["label"], t["nodeId"])
        return "%s [%s]" % (t["kind"], t["label"])

    def text(self, t, rnd):
        """A text nobody else could have written: this run's marker, the round, the slot."""
        if t["kind"] == "nodeLabel":        # one token: a flowchart label is word-wrapped
            return "%sR%d%s" % (self.marker, rnd, t["nodeId"])
        if t["kind"] == "behaviourDescription":
            return ["API test %s round %d: first bullet for %s" % (self.marker, rnd, t["label"]),
                    "API test %s round %d: second bullet for %s" % (self.marker, rnd, t["label"])]
        return "API test %s round %d: %s of %s" % (self.marker, rnd, t["kind"], t["label"])

    @staticmethod
    def pieces(t, text=None) -> list:
        """What to look for in a page or a Word file: a behaviour row prints each bullet."""
        text = t["current"] if text is None else text
        return [p for p in text.split("\n") if p.strip()] if t["kind"] == "behaviourDescription" \
            else [text]

    @staticmethod
    def _target(kind, row, **extra) -> dict:
        text = row.get("text") or ""
        t = {"kind": kind, "key": row.get("slotKey"),
             "label": row.get("label") or row.get("functionName") or row.get("slotKey"),
             "shownIn": list(row.get("shownIn") or []), "component": row.get("component"),
             "unit": row.get("unit"),
             "original": text,          # what the LLM wrote -- the text before any correction
             "current": text,           # what the document prints now, as far as this run knows
             "saved": [],               # every text saved, in order: what R5 must list
             "inForce": False, "undone": False, "undoneText": None, "slot": None,
             "functionId": row.get("functionId"), "externalCallerId": row.get("externalCallerId"),
             "row": row}                # as the catalog (R11, or R7 for a node) listed it
        t.update(extra)
        return t

    @staticmethod
    def _pick(rows, n) -> list:
        """Uncorrected slots, those a document prints and those with text first."""
        free = [r for r in rows if not r.get("isOverridden") and not r.get("isOrphaned")]
        free.sort(key=lambda r: (not r.get("shownIn"), not (r.get("text") or "").strip(),
                                 r.get("slotKey") or ""))
        return free[:n]

    # ---- saving, and checking what a save answered --------------------------------------------
    def _save_problems(self, t, slot, new) -> list:
        problems = ["missing %s" % f for f in SLOT_FIELDS + SAVE_FIELDS if f not in slot]
        llm = t["original"] if t["original"].strip() else None
        want = {"slotKind": t["kind"], "slotKey": t["key"], "text": new, "humanText": new,
                "llmText": llm, "isOverridden": True, "isOrphaned": False,
                "canUndo": llm is not None and llm.strip() != new.strip(),
                "previousText": t["current"] if t["current"].strip() else None,
                "firstEdit": not t["inForce"]}
        for k, v in want.items():
            if k in slot and slot[k] != v:
                problems.append("%s: want %r, got %r" % (k, v, slot[k]))
        if not slot.get("updatedAt"):
            problems.append("updatedAt is empty")
        if t["kind"] == "behaviourDescription" and slot.get("bullets") != new.split("\n"):
            problems.append("bullets: %r" % slot.get("bullets"))
        return problems

    def _saved(self, t, route, slot, new) -> None:
        problems = self._save_problems(t, slot, new)
        self.c.check("%s %s: saved, and answered as the slot now is" % (route, self.name(t)),
                     not problems, problems)
        t["saved"].append(new)
        t.update(current=new, inForce=True, slot=fields(slot))

    def save(self, t, text):
        """R3, or R6 for a behaviour row. Returns the answer, or None when it was refused."""
        if t["kind"] == "behaviourDescription":
            route, new = "R6", "\n".join(text)
            r = self.api.call("PUT", self.V("/overrides/behaviour"), json={
                "function_id": t["functionId"], "external_caller_id": t["externalCallerId"],
                "bullets": text})
        else:
            route, new = "R3", text
            r = self.api.call("PUT", self.V("/overrides/slot"), json={
                "slot_kind": t["kind"], "slot_key": t["key"], "text": text})
        if r.status_code != 200:
            self.c.check("%s %s: saved" % (route, self.name(t)), False,
                         (r.status_code, short(r.text)))
            return None
        slot = r.json()
        self._saved(t, route, slot, new)
        return slot

    def save_nodes(self, pairs):
        """R8: the labels of one flowchart, in one call -- only the ones that change."""
        flowchart_id = pairs[0][0]["flowchartId"]
        r = self.api.call("PUT", self.V("/flowcharts/labels"), json={
            "flowchart_id": flowchart_id,
            "labels": {t["nodeId"]: text for t, text in pairs}})
        what = "R8 %s: %d label(s) in one call" % (pairs[0][0]["label"], len(pairs))
        if r.status_code != 200:
            self.c.check(what, False, (r.status_code, short(r.text)))
            return None
        body = r.json()
        answered = [n.get("nodeId") for n in body.get("labels") or []]
        self.c.check(what + ": each saved node answered, in request order",
                     answered == [t["nodeId"] for t, _ in pairs], answered)
        for (t, text), slot in zip(pairs, body.get("labels") or []):
            self._saved(t, "R8", slot, text)
        dot = body.get("dot") or ""
        self.c.check("R8 answers the rebuilt diagram (dot), carrying the new label(s)",
                     all(text in dot for _, text in pairs), short(dot, 200))
        self.c.check("R8 says whether the Word picture is still owed (renderPending, renderJobs)",
                     isinstance(body.get("renderPending"), bool)
                     and isinstance(body.get("renderJobs"), list),
                     {k: body.get(k) for k in ("renderPending", "renderJobs")})
        return body

    # ---- what the documents show --------------------------------------------------------------
    def page_text(self) -> str:
        """What the document pages show: their JSON, and the flowchart pictures they point at.

        A page whose flowcharts are drawn on the server links each one as an SVG (`image_url`)
        instead of carrying its DOT, so a node label is in the picture, not in the JSON."""
        parts = []
        for d in self.documents():
            r = self.api.call("GET", self.P("/documents/%s/render" % d["id"]))
            if r.status_code != 200:
                continue
            body = r.json()
            parts.append(json.dumps(body, ensure_ascii=False))
            for url in _svg_urls(body):
                pic = self.api.call("GET", "/" + url.lstrip("/"))
                if pic.status_code == 200:
                    parts.append(pic.text)
        return norm(" ".join(parts))

    def check_page(self, kinds, when):
        want = [(t, p) for t in self.targets if t["inForce"] and t["shownIn"] and t["kind"] in kinds
                for p in self.pieces(t)]
        if not want:
            self.c.skip("the document page shows the corrections %s" % when,
                        "no corrected slot of %s is printed by a document of this version"
                        % "/".join(kinds))
            return
        page = self.page_text()
        missing = [(self.name(t), p) for t, p in want if norm(p) not in page]
        self.c.check("the document page shows %d corrected text(s) %s" % (len(want), when),
                     not missing, missing[:6])

    # ---- the steps ----------------------------------------------------------------------------
    def run(self) -> bool:
        steps = [
            ("1. Sign in", self.step_signin, True),
            ("2. Onboard the project", self.step_onboard, True),
            ("3. Generate a version", self.step_generate, True),
            ("4. The version before any correction", self.step_before, True),
            ("5. First corrections, every kind", self.step_first, True),
            ("6. Every route reads the same corrections", self.step_readback, False),
            ("7. Correct again", self.step_second, False),
            ("8. Two saves of one slot at the same moment", self.step_concurrent, False),
            ("9. Undo", self.step_undo, False),
            ("10. Mistakes a client can make", self.step_negative, False),
            ("11. Re-export and the Word file", self.step_reexport, False),
            ("12. Correct again after the export", self.step_after_export, False),
            ("13. The next version carries the corrections", self.step_carry, False),
        ]
        for title, step, required in steps:
            self.c.begin(title)
            try:
                step()
            except StopTest as exc:
                self.c.check("this step could not run", False, str(exc))
                if required:
                    return False
            except Exception as exc:               # noqa: BLE001 -- report it, go on
                self.c.check("this step raised %s" % type(exc).__name__, False, repr(exc))
                if required:
                    return False
        return True

    def step_signin(self):
        self.api.signin()
        self.c.check("POST /auth/signin answers an access token", bool(self.api.token))
        r = self.api.call("GET", "/projects", auth=False)
        self.c.check("a request without a token is refused (401)", r.status_code == 401,
                     r.status_code)

    def step_onboard(self):
        if self.pid:
            r = self.api.call("GET", self.P())
            if r.status_code != 200:
                raise StopTest("project %s: %s %s" % (self.pid, r.status_code, short(r.text)))
            self.c.check("the existing project %s is reachable" % self.pid, True)
        else:
            create = fill((self.cfg.get("onboard") or {}).get("create_project") or {}, self.values)
            if not create:
                raise StopTest("onboard.create_project is empty in the config")
            r = self.api.call("POST", "/projects", json=create)
            ok = self.c.check("POST /projects creates the project", r.status_code in (200, 201),
                              (r.status_code, short(r.text)))
            if not ok:
                raise StopTest("the project could not be created")
            self.pid = r.json()["project"]["id"]
            self.c.note("project id: %s" % self.pid)
        project = body_of(self.api.call("GET", self.P())).get("project") or {}
        self.ids["projectId"] = self.pid
        self.values["projectId"] = self.pid
        self.c.check("you are an admin of the project (needed to start a job and a re-export)",
                     project.get("my_role") == "admin", project.get("my_role"))
        self.c.check("the project has an architecture (needed to start a job)",
                     bool(project.get("architecture_layers")),
                     "a project onboarded with analyzer.py keeps its layers in its workspace",
                     warn=bool(self.vid))

    def latest_commit(self, branch=None) -> str:
        r = self.api.call("GET", self.P("/commits"), params={"per_page": 100})
        commits = body_of(r).get("commits") or []
        if branch:
            commits = [c for c in commits if c.get("branch") in (branch, None)] or commits
        if not commits:
            raise StopTest("the server lists no commits for this project's repository (%s %s) "
                           "-- set generate.start_job.commit_sha" % (r.status_code, short(r.text)))
        return max(commits, key=lambda c: c.get("committed_at") or "")["sha"]

    def step_generate(self):
        gen = self.cfg.get("generate") or {}
        if self.vid:
            r = self.api.call("GET", self.P("/versions/%s" % self.vid))
            if r.status_code != 200:
                raise StopTest("version %s: %s %s" % (self.vid, r.status_code, short(r.text)))
            version = body_of(r).get("version") or body_of(r)
            self.c.check("the existing version %s is reachable" % self.vid, True)
            self.values["versionTag"] = version.get("tag") or version.get("version_tag") or ""
        else:
            body = fill(gen.get("start_job") or {}, self.values)
            sha = str(body.get("commit_sha") or "").strip()
            if not sha or sha.lower() == "latest" or sha.startswith("<"):
                sha = self.latest_commit(gen.get("branch"))
                self.c.note("commit: %s (the newest the server lists)" % sha)
            body["commit_sha"] = sha
            self.job_body = body
            r = self.api.call("POST", self.P("/jobs"), json=body)
            ok = self.c.check("POST /projects/{projectId}/jobs starts the generation",
                              r.status_code in (200, 201, 202), (r.status_code, short(r.text)))
            if not ok:
                raise StopTest("the job could not be started")
            job, seconds = self.wait_job(r.json()["job_id"], "generation")
            self.ids["generationJob"] = job["id"]
            ok = self.c.check("the generation completes (%d s)" % seconds,
                              job["status"] == "complete", job.get("error_message"))
            if not ok:
                raise StopTest("the generation did not complete: %s" % job.get("error_message"))
            self.vid = job["version_id"]
            self.values["versionTag"] = job.get("version_tag") or body.get("version_tag") or ""
            self.c.note("version id: %s (tag %s)" % (self.vid, self.values["versionTag"]))
        self.values["versionId"] = self.vid
        self.ids["versionId"] = self.vid
        docs = self.documents()
        self.c.check("the version has documents (%d)" % len(docs), bool(docs),
                     "GET /projects/{projectId}/documents?version_id= answered none")

    def step_before(self):
        if self.fresh:
            r9 = self.readiness()
            self.c.check("R9: nothing is stale before any correction",
                         r9.get("stale") is False and r9.get("overrideCount") == 0, r9)
            r1 = body_of(self.api.call("GET", self.V("/overrides"), params={"limit": 1000}))
            self.c.check("R1: no corrections yet", r1.get("total") == 0, r1.get("total"))
        n = int(self.opt["slots_per_kind"])
        for kind in self.opt["kinds"]:
            rows = self.slots(kind)
            self.catalog[kind] = rows
            if kind == "nodeLabel":
                self.c.check("R11 nodeLabel: %d flowcharts listed, each with its flowchartId"
                             % len(rows), all(r.get("flowchartId") for r in rows))
                continue
            no_shape = [r.get("slotKey") for r in rows if any(f not in r for f in SLOT_FIELDS)]
            wrong = [r.get("slotKey") for r in rows
                     if not r.get("isOverridden") and not r.get("isOrphaned")
                     and (r.get("humanText") is not None or r.get("canUndo")
                          or r.get("llmText") != ((r.get("text") or "").strip() and r.get("text")
                                                  or None))]
            self.c.check("R11 %s: %d slots, each a Slot; an uncorrected one has llmText = text and "
                         "no undo" % (kind, len(rows)), not no_shape and not wrong,
                         {"withoutSlotFields": no_shape[:4], "wrong": wrong[:4]})
            for row in self._pick(rows, n):
                self.targets.append(self._target(kind, row))
        if "nodeLabel" in self.opt["kinds"]:
            self._pick_flowchart()
        for kind in self.opt["kinds"]:
            chosen = self.by_kind(kind)
            if not chosen:
                self.c.skip("a %s to correct" % kind, "this version has none")
                continue
            for t in chosen:
                self.c.note("will correct %s -- printed in %s" % (
                    self.name(t), ", ".join(t["shownIn"]) or "no document of this version"))
        if not self.targets:
            raise StopTest("nothing to correct in this version")
        # R2 reads a slot nobody corrected -- the same slot R11 (or R7) lists.
        seen, wrong = set(), []
        for t in self.targets:
            if t["kind"] in seen:
                continue
            seen.add(t["kind"])
            r = self.read(t)
            listed = t.get("row")
            if r.status_code != 200 or (listed and diff(fields(listed), fields(r.json()))):
                wrong.append((self.name(t), r.status_code,
                              diff(fields(listed), fields(body_of(r))) if listed else ""))
        self.c.check("R2 reads a slot nobody corrected (one per kind): 200, as R11/R7 list it",
                     not wrong, wrong[:4])

    def _pick_flowchart(self):
        flowcharts = sorted(self.catalog.get("nodeLabel") or [],
                            key=lambda f: (not f.get("shownIn"), -(f.get("nodeCount") or 0)))
        for fc in flowcharts[:8]:
            body = self.labels(fc["flowchartId"])
            nodes = [n for n in body.get("labels") or []
                     if (n.get("text") or "").strip() and not n.get("isOverridden")
                     and not n.get("isOrphaned")]
            if len(nodes) < 2:
                continue
            no_shape = [n.get("nodeId") for n in body["labels"]
                        if any(f not in n for f in SLOT_FIELDS + ("flowchartId", "nodeId"))]
            self.c.check("R7 %s: %d nodes, each a Slot with its flowchartId and nodeId"
                         % (fc.get("functionName"), len(body["labels"])), not no_shape, no_shape)
            self.c.check("R7 carries the diagram (dot)", bool(body.get("dot")))
            self.flowchart = fc
            for node in nodes[:2]:
                self.targets.append(self._target(
                    "nodeLabel", node, label=fc.get("functionName") or fc["flowchartId"],
                    shownIn=list(fc.get("shownIn") or []), unit=fc.get("unit"),
                    component=fc.get("component"), flowchartId=fc["flowchartId"],
                    nodeId=node["nodeId"]))
            return

    def step_first(self):
        # Descriptions first: R10 must hold what they queued before any other save can take a
        # queued slot off the list (a corrected slot owes nothing -- REQ-CS-03).
        queued = []
        for t in self.by_kind("description"):
            slot = self.save(t, self.text(t, 1))
            for q in (slot or {}).get("queuedForRegeneration") or []:
                queued.append((q.get("slotKind"), q.get("slotKey")))
        if self.by_kind("description"):
            pending = self.queued()
            self.c.check("R10 lists everything the description saves queued (%d)" % len(queued),
                         all(q in pending for q in queued), [q for q in queued if q not in pending])
        for kind in R3_KINDS[1:] + ("behaviourDescription",):
            for t in self.by_kind(kind):
                self.save(t, self.text(t, 1))
        nodes = self.by_kind("nodeLabel")
        if nodes:
            self.save_nodes([(t, self.text(t, 1)) for t in nodes])
        if not any(t["inForce"] for t in self.targets):
            raise StopTest("no correction could be saved")

    def step_readback(self):
        done = [t for t in self.targets if t["inForce"]]
        wrong = []
        for t in done:
            r = self.read(t)
            if r.status_code != 200 or diff(t["slot"], fields(r.json())):
                wrong.append((self.name(t), r.status_code, diff(t["slot"], fields(body_of(r)))))
        self.c.check("R2 reads each corrected slot as its save answered (%d)" % len(done),
                     not wrong, wrong[:4])
        r1 = body_of(self.api.call("GET", self.V("/overrides"), params={"limit": 1000}))
        listed = {(i.get("slotKind"), i.get("slotKey")): i for i in r1.get("overrides") or []}
        wrong = [(self.name(t), diff(t["slot"], fields(listed.get((t["kind"], t["key"])))))
                 for t in done if diff(t["slot"], fields(listed.get((t["kind"], t["key"]))))]
        self.c.check("R1 lists every correction, each as its save answered", not wrong, wrong[:4])
        if self.fresh:
            self.c.check("R1 total = %d, the corrections this run made" % len(done),
                         r1.get("total") == len(done), r1.get("total"))
        wrong = []
        for kind in sorted({t["kind"] for t in done} - {"nodeLabel"}):
            for t in [x for x in done if x["kind"] == kind]:
                rows = self.slots(kind, unit=t["unit"]) if t.get("unit") else self.slots(kind)
                row = next((x for x in rows if x.get("slotKey") == t["key"]), None)
                if row is None or diff(t["slot"], fields(row)):
                    wrong.append((self.name(t), diff(t["slot"], fields(row)) if row else "absent"))
        self.c.check("R11 lists each corrected slot as its save answered", not wrong, wrong[:4])
        nodes = [t for t in done if t["kind"] == "nodeLabel"]
        if nodes:
            by_id = {n.get("nodeId"): n for n in self.labels(nodes[0]["flowchartId"]).get("labels")
                     or []}
            wrong = [(self.name(t), diff(t["slot"], fields(by_id.get(t["nodeId"]))))
                     for t in nodes if diff(t["slot"], fields(by_id.get(t["nodeId"])))]
            self.c.check("R7 shows each corrected node as R8 answered it", not wrong, wrong)
        wrong = [(self.name(t), self.history(t), t["saved"]) for t in done
                 if self.history(t) != t["saved"]]
        self.c.check("R5 lists each slot's edits, oldest first", not wrong, wrong[:3])
        r9 = self.readiness()
        self.c.check("R9: stale -- the corrections are not in the Word file yet",
                     r9.get("stale") is True, r9)
        self.c.check("R9 counts the corrections (%d)" % len(done),
                     (r9.get("overrideCount") == len(done)) if self.fresh
                     else (r9.get("overrideCount") or 0) >= len(done), r9.get("overrideCount"))
        pending = self.queued()
        owed = [self.name(t) for t in done if (t["kind"], t["key"]) in pending]
        self.c.check("R10: no corrected slot is queued for regeneration", not owed, owed)
        if self.opt["check_page"]:
            self.check_page(PAGE_AFTER_SAVE, "on the next page load")

    def step_second(self):
        for kind in R3_KINDS + ("behaviourDescription",):
            for t in [x for x in self.by_kind(kind) if x["inForce"]]:
                self.save(t, self.text(t, 2))
        nodes = [t for t in self.by_kind("nodeLabel") if t["inForce"]][:1]
        if nodes:
            self.save_nodes([(nodes[0], self.text(nodes[0], 2))])
        done = [t for t in self.targets if len(t["saved"]) >= 2]
        wrong = [(self.name(t), self.history(t)) for t in done if self.history(t) != t["saved"]]
        self.c.check("R5 lists both edits of each slot, oldest first (%d slots)" % len(done),
                     not wrong, wrong[:3])

    def step_concurrent(self):
        if not self.opt["concurrent_saves"]:
            self.c.skip("two saves at once", "test.concurrent_saves is false")
            return
        t = next((x for x in self.targets if x["inForce"] and x["kind"] in R3_KINDS), None)
        if t is None:
            self.c.skip("two saves at once", "no corrected slot to save twice")
            return
        texts = ["API test %s: concurrent save %s of %s" % (self.marker, s, t["label"])
                 for s in ("A", "B")]
        before = t["current"]

        def put(api, text):
            return api.call("PUT", self.V("/overrides/slot"),
                            json={"slot_kind": t["kind"], "slot_key": t["key"], "text": text})

        with concurrent.futures.ThreadPoolExecutor(2) as pool:
            answers = list(pool.map(put, [self.api.twin(), self.api.twin()], texts))
        codes = [a.status_code for a in answers]
        sqlite = self.database == "sqlite"
        self.c.note("%s: two saves sent together; answered %s" % (self.name(t), codes))
        both = self.c.check("both saves answer 200 -- the second waits for the first",
                            codes == [200, 200], codes, warn=sqlite)
        now = body_of(self.read(t))
        final = now.get("text")
        self.c.check("the slot holds one of the two texts: the last save wins", final in texts,
                     final)
        if final in texts:
            winner = final
            loser = texts[1 - texts.index(final)]
            if both:
                won = answers[texts.index(winner)].json()
                self.c.check("the winning save's previousText is the other save's text",
                             won.get("previousText") == loser, won.get("previousText"),
                             warn=sqlite)
                history = self.history(t)
                self.c.check("R5 keeps both saves, the winner last", history[-2:] == [loser, winner],
                             history[-3:], warn=sqlite)
                t["saved"] += [loser, winner]
            else:
                t["saved"].append(winner)
            t.update(current=winner, slot=fields(now))
        elif final == before:
            t["slot"] = fields(now)

    def step_undo(self):
        # Per kind: one slot the LLM wrote text for (undo goes back to it) and one it left empty
        # (undo is refused: nothing to go back to) -- whichever of the two this version has.
        chosen = []
        for kind in ALL_KINDS:
            live = [t for t in self.by_kind(kind) if t["inForce"]]
            chosen += [x for x in (next((t for t in live if t["original"].strip()), None),
                                   next((t for t in live if not t["original"].strip()), None))
                       if x is not None]
        for t in chosen:
            r = self.api.call("DELETE", self.V("/overrides/slot"),
                              params={"slot_kind": t["kind"], "slot_key": t["key"]})
            if not t["original"].strip():
                self.c.check("R4 %s: refused (409) -- the LLM wrote nothing to go back to"
                             % self.name(t), r.status_code == 409, (r.status_code, short(r.text)))
                continue
            if r.status_code != 200:
                self.c.check("R4 %s: undone" % self.name(t), False, (r.status_code, short(r.text)))
                continue
            body = r.json()
            problems = ["missing %s" % f for f in SLOT_FIELDS + SAVE_FIELDS if f not in body]
            want = {"text": t["original"], "llmText": t["original"], "humanText": t["original"],
                    "isOverridden": True, "isOrphaned": False, "canUndo": False,
                    "previousText": t["current"], "firstEdit": False}
            problems += ["%s: want %r, got %r" % (k, v, body.get(k))
                         for k, v in want.items() if body.get(k) != v]
            if t["kind"] == "nodeLabel":
                if t["current"] in (body.get("dot") or ""):
                    problems.append("dot still carries %r" % t["current"])
                if "renderPending" not in body:
                    problems.append("no renderPending")
            self.c.check("R4 %s: back to the LLM's text, answered like a save" % self.name(t),
                         not problems, problems)
            t["saved"].append(t["original"])
            t.update(undoneText=t["current"], current=t["original"], undone=True,
                     slot=fields(body))
        undone = [t for t in self.targets if t["undone"]]
        if undone:
            t = undone[0]
            r = self.api.call("DELETE", self.V("/overrides/slot"),
                              params={"slot_kind": t["kind"], "slot_key": t["key"]})
            self.c.check("a second undo of %s changes nothing (200, same text)" % self.name(t),
                         r.status_code == 200 and body_of(r).get("text") == t["original"],
                         (r.status_code, body_of(r).get("text")))
            if r.status_code == 200:
                t["saved"].append(t["original"])
                t["slot"] = fields(body_of(r))
            wrong = [(self.name(x), diff(x["slot"], fields(body_of(self.read(x))))) for x in undone
                     if diff(x["slot"], fields(body_of(self.read(x))))]
            self.c.check("R2 reads each undone slot as the undo answered it", not wrong, wrong)
            wrong = [(self.name(x), self.history(x)) for x in undone
                     if self.history(x) != x["saved"]]
            self.c.check("R5 records each undo as an edit back to the LLM's text", not wrong,
                         wrong[:3])
        spare = next((r for kind in R3_KINDS for r in self.catalog.get(kind) or []
                      if not r.get("isOverridden") and not r.get("isOrphaned")
                      and all(r.get("slotKey") != t["key"] for t in self.targets)), None)
        if spare:
            r = self.api.call("DELETE", self.V("/overrides/slot"),
                              params={"slot_kind": spare["slotKind"], "slot_key": spare["slotKey"]})
            self.c.check("R4 on a slot nobody corrected is refused (409)", r.status_code == 409,
                         (r.status_code, short(r.text)))

    def step_negative(self):
        if not self.opt["negative_checks"]:
            self.c.skip("the mistakes a client can make", "test.negative_checks is false")
            return
        text_t = next(iter(self.by_kind(*R3_KINDS)), None)
        node_t = next(iter(self.by_kind("nodeLabel")), None)
        beh_t = next(iter(self.by_kind("behaviourDescription")), None)
        cases = []
        if text_t:
            k, key = text_t["kind"], text_t["key"]
            cases += [
                ("R3 with empty text is 422", "PUT", "/overrides/slot",
                 {"json": {"slot_kind": k, "slot_key": key, "text": ""}}, 422),
                ("R3 with whitespace-only text is 422", "PUT", "/overrides/slot",
                 {"json": {"slot_kind": k, "slot_key": key, "text": "   "}}, 422),
                ("R3 in camelCase (slotKind) is 422 -- requests are snake_case", "PUT",
                 "/overrides/slot", {"json": {"slotKind": k, "slotKey": key, "text": "x"}}, 422),
                ("R3 on a slot that is not in the version is 404", "PUT", "/overrides/slot",
                 {"json": {"slot_kind": "description", "slot_key": "No|Such|function|",
                           "text": "x"}}, 404),
                ("R3 with a malformed key is 400", "PUT", "/overrides/slot",
                 {"json": {"slot_kind": "description", "slot_key": "one\u0001two", "text": "x"}},
                 400),
                ("R2 on a slot that is not in the version is 404", "GET", "/overrides/slot",
                 {"params": {"slot_kind": "description", "slot_key": "No|Such|function|"}}, 404),
            ]
        if node_t:
            fid = node_t["flowchartId"]
            cases += [
                ("R8 with no labels is 400", "PUT", "/flowcharts/labels",
                 {"json": {"flowchart_id": fid, "labels": {}}}, 400),
                ("R8 with an empty label is 422", "PUT", "/flowcharts/labels",
                 {"json": {"flowchart_id": fid, "labels": {node_t["nodeId"]: "  "}}}, 422),
                ("R8 with one node's key as the flowchart id is 400", "PUT", "/flowcharts/labels",
                 {"json": {"flowchart_id": node_t["key"], "labels": {node_t["nodeId"]: "x"}}}, 400),
                ("R7 on a flowchart that is not in the version is 404", "GET", "/flowcharts/labels",
                 {"params": {"flowchart_id": "No|Such|function|"}}, 404),
            ]
        if beh_t:
            cases += [
                ("R6 with no bullets is 422", "PUT", "/overrides/behaviour",
                 {"json": {"function_id": beh_t["functionId"],
                           "external_caller_id": beh_t["externalCallerId"], "bullets": []}}, 422),
                ("R6 on a call that is not in the version is 404", "PUT", "/overrides/behaviour",
                 {"json": {"function_id": beh_t["functionId"],
                           "external_caller_id": "No|Such|caller|", "bullets": ["x"]}}, 404),
            ]
        cases += [
            ("R11 with an unknown slot kind is 422", "GET", "/slots",
             {"params": {"slot_kind": "notAKind"}}, 422),
        ]
        for name, method, path, kwargs, code in cases:
            r = self.api.call(method, self.V(path), **kwargs)
            self.c.check(name, r.status_code == code, (r.status_code, short(r.text)))

        # R3 offers all seven kinds but saves five: a behaviour row and a flowchart label have
        # routes of their own, and the 501 must say which -- its path, with the ids from the key.
        for t, route, path in ((beh_t, "R6", "/overrides/behaviour"),
                               (node_t, "R8", "/flowcharts/labels")):
            if not t:
                continue
            r = self.api.call("PUT", self.V("/overrides/slot"), json={
                "slot_kind": t["kind"], "slot_key": t["key"], "text": "x"})
            detail = str(body_of(r).get("detail") or "")
            self.c.check("R3 cannot save a %s (501) and names %s, its path and the ids"
                         % (t["kind"], route),
                         r.status_code == 501 and ("%s: PUT /api/v1%s" % (route, self.V(path)))
                         in detail and repr(t.get("functionId") or t.get("flowchartId")) in detail,
                         (r.status_code, short(detail)))

        # A NUL character: PostgreSQL cannot store one in text or JSONB, so a save carrying it
        # was a 500 there -- and SQLite stored it. Each save refuses it (422) and stores nothing.
        nul = []
        if text_t:
            nul.append(("R3 with a NUL character in the text", text_t, "/overrides/slot",
                        {"slot_kind": text_t["kind"], "slot_key": text_t["key"],
                         "text": "API test %s NUL \u0000 here" % self.marker}))
        if beh_t:
            nul.append(("R6 with a NUL character in a bullet", beh_t, "/overrides/behaviour",
                        {"function_id": beh_t["functionId"],
                         "external_caller_id": beh_t["externalCallerId"],
                         "bullets": ["API test %s NUL \u0000 here" % self.marker]}))
        if node_t:
            nul.append(("R8 with a NUL character in a label", node_t, "/flowcharts/labels",
                        {"flowchart_id": node_t["flowchartId"],
                         "labels": {node_t["nodeId"]: "%sNUL\u0000" % self.marker}}))
        for name, t, path, body in nul:
            before = body_of(self.read(t))
            r = self.api.call("PUT", self.V(path), json=body)
            after = body_of(self.read(t))
            self.c.check(name + " is 422 and saves nothing",
                         r.status_code == 422 and after == before, (r.status_code, short(r.text)))

        r = self.api.call("GET", "/projects/%s/versions/verNOSUCH/overrides" % self.pid)
        self.c.check("a version that is not the project's is 404", r.status_code == 404,
                     (r.status_code, short(r.text)))
        if node_t:
            before = {n.get("nodeId"): n.get("text")
                      for n in self.labels(node_t["flowchartId"]).get("labels") or []}
            r = self.api.call("PUT", self.V("/flowcharts/labels"), json={
                "flowchart_id": node_t["flowchartId"],
                "labels": {node_t["nodeId"]: "%sMUSTNOTSAVE" % self.marker, "nNOSUCH": "x"}})
            after = {n.get("nodeId"): n.get("text")
                     for n in self.labels(node_t["flowchartId"]).get("labels") or []}
            self.c.check("R8 naming a node that does not exist is 404 and saves nothing -- not "
                         "even the good label", r.status_code == 404 and after == before,
                         (r.status_code, short(r.text)))

    def step_reexport(self):
        if not self.opt["reexport"]:
            self.c.skip("the re-export", "test.reexport is false")
            return
        r9 = self.readiness()
        self.c.check("R9 before: stale -- the Word file does not carry the corrections yet",
                     r9.get("stale") is True, r9)
        r = self.api.call("POST", self.P("/versions/%s/reexport" % self.vid))
        detail = body_of(r).get("detail")
        if r.status_code == 409 and isinstance(detail, dict) and detail.get("job_id"):
            job_id = detail["job_id"]
            self.c.note("a re-export was already running; following %s" % job_id)
        else:
            ok = self.c.check("POST /versions/{versionId}/reexport starts a re-export job (202)",
                              r.status_code == 202, (r.status_code, short(r.text)))
            if not ok:
                return
            job_id = r.json()["job_id"]
        job, seconds = self.wait_job(job_id, "re-export")
        self.ids["reexportJob"] = job_id
        if not self.c.check("the re-export completes (%d s)" % seconds,
                            job["status"] == "complete", job.get("error_message")):
            return
        r9 = self.readiness()
        self.c.check("R9 after: nothing stale, no picture owed",
                     r9.get("stale") is False and r9.get("pendingRenders") == 0, r9)
        self.c.check("R9 names that re-export, complete",
                     (r9.get("reexport") or {}).get("jobId") == job_id
                     and (r9.get("reexport") or {}).get("status") == "complete", r9.get("reexport"))
        if self.opt["check_word"]:
            self.check_word()
        if self.opt["check_page"]:
            self.check_page(("structDescription",), "after the re-export (struct descriptions)")

    def check_word(self):
        import importlib.util
        if importlib.util.find_spec("docx") is None:
            self.c.skip("the Word file", "python-docx is not installed (pip install python-docx)")
            return
        docs, text, bad = self.documents(), [], []
        for d in docs:
            r = self.api.call("GET", self.P("/documents/%s/download" % d["id"]))
            if r.status_code != 200 or not zipfile.is_zipfile(io.BytesIO(r.content)):
                bad.append((d.get("name"), r.status_code, len(r.content)))
                continue
            text.append(docx_text(r.content))
        self.c.check("every document downloads as a Word file (%d)" % len(docs), not bad, bad)
        word = norm(" ".join(text))
        want = [(t, p) for t in self.targets if t["inForce"] and t["shownIn"]
                and t["kind"] != "nodeLabel" for p in self.pieces(t)]
        missing = [(self.name(t), p) for t, p in want if norm(p) not in word]
        self.c.check("the Word file carries the %d corrected text(s) a document prints "
                     "(node labels are in the pictures)" % len(want), not missing, missing[:6])
        gone = [(self.name(t), p) for t in self.targets if t["undone"] and t["shownIn"]
                and t["kind"] != "nodeLabel" for p in self.pieces(t, t["undoneText"])
                if norm(p) in word]
        self.c.check("the Word file no longer carries the undone corrections", not gone, gone[:6])

    def step_after_export(self):
        t = next((x for x in self.targets if x["undone"] and x["kind"] in R3_KINDS), None) \
            or next((x for x in self.targets if x["undone"]), None)
        if t is None:
            self.c.skip("a correction after the export", "no undone slot to correct again")
            return
        if t["kind"] == "nodeLabel":
            self.save_nodes([(t, self.text(t, 3))])
        else:
            self.save(t, self.text(t, 3))
        t["undone"] = False
        r9 = self.readiness()
        self.c.check("R9: stale again -- that correction is not in the Word file yet",
                     r9.get("stale") is True, r9)

    def step_carry(self):
        if not self.opt["carry_forward"]:
            self.c.skip("the next version", "test.carry_forward is false")
            return
        if not self.job_body:
            self.c.skip("the next version", "this run did not generate the version, so it has no "
                                            "job body to generate the next one from")
            return
        body = copy.deepcopy(self.job_body)
        body.update(version_tag="%s-next" % self.values.get("versionTag", self.marker),
                    reference_version_id=self.vid, mode="auto")
        r = self.api.call("POST", self.P("/jobs"), json=body)
        if not self.c.check("POST /projects/{projectId}/jobs starts the next version from this "
                            "one (reference_version_id)", r.status_code in (200, 201, 202),
                            (r.status_code, short(r.text))):
            return
        job, seconds = self.wait_job(r.json()["job_id"], "next version")
        self.ids["nextVersionJob"] = job["id"]
        if not self.c.check("the next version completes (%d s)" % seconds,
                            job["status"] == "complete", job.get("error_message")):
            return
        v2 = job["version_id"]
        self.ids["nextVersionId"] = v2
        self.c.note("next version id: %s" % v2)
        version = body_of(self.api.call("GET", self.P("/versions/%s" % v2))).get("version") or {}
        self.c.check("it is generated incrementally from this version",
                     version.get("decision") == "incremental"
                     and version.get("baseline_version_id") == self.vid,
                     {k: version.get(k) for k in ("decision", "baseline_version_id")})
        r1 = body_of(self.api.call("GET", self.V("/overrides", v2), params={"limit": 1000}))
        listed = {(i.get("slotKind"), i.get("slotKey")): i for i in r1.get("overrides") or []}
        wrong = []
        for t in [x for x in self.targets if x["inForce"]]:
            got = listed.get((t["kind"], t["key"]))
            want = {"text": t["current"], "humanText": t["current"], "isOverridden": True,
                    "isOrphaned": False}
            bad = ["%s: want %r, got %r" % (k, v, (got or {}).get(k))
                   for k, v in want.items() if (got or {}).get(k) != v]
            if bad:
                wrong.append((self.name(t), bad if got else "absent"))
        self.c.check("every correction is carried into the next version, in force and printed "
                     "(the code did not change)", not wrong, wrong[:4])
        self.c.note("regeneration queue: this version %d, the next %d"
                    % (len(self.queued()), len(self.queued(v2))))

    # ---- the report ---------------------------------------------------------------------------
    def report(self, seconds) -> dict:
        return {
            "finishedAt": dt.datetime.now().isoformat(timespec="seconds"),
            "seconds": seconds,
            "server": self.api.base,
            "marker": self.marker,
            "ids": self.ids,
            "counts": self.c.counts(),
            "checks": self.c.items,
            "targets": [{k: t.get(k) for k in ("kind", "key", "label", "shownIn", "original",
                                               "saved", "current", "undone")}
                        for t in self.targets],
        }


# ---------------------------------------------------------------------------
# entry point
# ---------------------------------------------------------------------------
def make_sample_repo(dest) -> int:
    """A git repository holding SampleCppProject, for a server on this machine to clone."""
    src = os.path.join(REPO_ROOT, "SampleCppProject")
    dest = os.path.abspath(dest)
    if os.path.isdir(dest) and os.listdir(dest):
        print("%s exists and is not empty" % dest, file=sys.stderr)
        return 2
    shutil.copytree(src, dest, dirs_exist_ok=True)

    def git(*args):
        return subprocess.run(["git", "-C", dest, *args], capture_output=True, text=True,
                              check=True, shell=(os.name == "nt"))
    git("init", "-q")
    git("symbolic-ref", "HEAD", "refs/heads/main")
    git("add", "-A")
    git("-c", "user.email=review-api-test@example.com", "-c", "user.name=review-api-test",
        "commit", "-q", "-m", "SampleCppProject")
    sha = git("rev-parse", "HEAD").stdout.strip()
    print("sample repository: %s\ncommit:            %s\n\nIn your config set "
          "onboard.create_project.repo_url to \"%s\"." % (dest, sha, dest.replace("\\", "/")))
    return 0


def in_process_transport():
    """The API inside this process -- no server needed; the database is DATABASE_URL's."""
    sys.path[:0] = [REPO_ROOT, os.path.join(REPO_ROOT, "engine")]
    from fastapi.testclient import TestClient
    from api.db.postgres.database import SqlDatabase
    from api.db.session import get_db
    from api.main import app
    from core.db import get_engine
    app.dependency_overrides[get_db] = lambda: SqlDatabase(get_engine())
    client = TestClient(app)
    client.__enter__()                            # the startup events
    return client


def main(argv=None) -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(errors="replace")
    ap = argparse.ArgumentParser(
        description="Test the whole review & update feature through the REST API.")
    ap.add_argument("--config", default=os.path.join(HERE, "config.json"),
                    help="the JSON config (see config.example.json). Default: config.json here")
    ap.add_argument("--project-id", help="use this existing project instead of creating one")
    ap.add_argument("--version-id", help="use this existing version instead of generating one")
    ap.add_argument("--make-sample-repo", metavar="DIR",
                    help="create a git copy of SampleCppProject in DIR for the server to clone, "
                         "then stop")
    ap.add_argument("--in-process", action="store_true",
                    help="run the API inside this process instead of calling a server "
                         "(developers; needs DATABASE_URL or config.local.json's db)")
    args = ap.parse_args(argv)
    if args.make_sample_repo:
        return make_sample_repo(args.make_sample_repo)
    try:
        with open(args.config, encoding="utf-8") as fh:
            cfg = json.load(fh)
    except (OSError, ValueError) as exc:
        print("cannot read the config %s: %s" % (args.config, exc), file=sys.stderr)
        return 2
    server = cfg.get("server") or {}
    transport = in_process_transport() if args.in_process else None
    api = Api("http://testserver" if args.in_process else server.get("base_url"),
              server.get("email"), server.get("password"),
              timeout=float(server.get("request_timeout_seconds") or 300),
              verify=server.get("verify_tls", True), transport=transport)
    checks = Checks()
    print("Review & Update API test -- %s as %s" % (api.base, server.get("email")), flush=True)
    started = time.time()
    try:
        test = ReviewApiTest(api, cfg, checks, args)
    except StopTest as exc:
        print("cannot start: %s" % exc, file=sys.stderr)
        return 2
    completed = test.run()
    seconds = int(time.time() - started)
    counts = checks.counts()
    report = test.report(seconds)
    path = test.opt.get("report_file") or DEFAULT_OPTIONS["report_file"]
    try:
        with open(path, "w", encoding="utf-8") as fh:
            json.dump(report, fh, indent=2, ensure_ascii=False)
    except OSError as exc:
        print("could not write the report %s: %s" % (path, exc), file=sys.stderr)
    print("\n== Summary (%d s)" % seconds)
    print("   PASS %(PASS)d   FAIL %(FAIL)d   WARN %(WARN)d   SKIP %(SKIP)d" % counts)
    for item in checks.items:
        if item["result"] == "FAIL":
            print("   FAIL  [%s] %s" % (item["step"], item["check"]))
    print("   ids: %s" % json.dumps(test.ids))
    print("   report: %s" % os.path.abspath(path))
    if not completed:
        return 2
    return 1 if counts["FAIL"] else 0


if __name__ == "__main__":
    sys.exit(main())
