"""Comparing the Word files of two versions -- how the REST test tools show that two ways of
making a version give the same documents: paused and resumed against never paused, killed and
resumed, incremental against full (tools/incremental_test uses it too).

Each pair of documents with the same name is looked at twice:

* what it says -- tools/doccheck reads both into what they describe (headings, units, the
  interface rows and their cells, behaviour diagrams, test cases) and compares them on five
  levels. A P1 or P2 finding is a difference in what the document says; wording alone is P4, and
  the second look's.
* word for word -- the whole text, paragraphs and table cells, with the versions' names and ids,
  and dates, taken out. Without the LLM every word comes from the code; with it, every LLM text
  comes from a cache the project's versions share, so a difference is a text written again.

The documents are kept under the run's log folder, `documents/<version id>/`, to open by hand.
"""
from __future__ import annotations

import io
import os
import re
import sys
import zipfile

from jst_core import REPO_ROOT, short

_DATE = re.compile(r"\d{4}-\d{2}-\d{2}([ T]\d{2}:\d{2}(:\d{2})?)?")


def _doccheck():
    tools = os.path.join(REPO_ROOT, "tools")
    if tools not in sys.path:
        sys.path.insert(0, tools)
    from doccheck import blocks, compare, match, rules, swe3, swe4
    return blocks, compare, match, rules, swe3, swe4


def fetch(ctx, key, vid) -> dict:
    """{document name ("SWE.3 Lib"): its .docx, downloaded into the log folder} of version `vid`."""
    folder = os.path.join(ctx.log_dir, "documents", vid)
    os.makedirs(folder, exist_ok=True)
    out = {}
    for d in ctx.documents(key, vid):
        r = ctx.admin.call("GET", ctx.P(key, "/documents/%s/download" % d["id"]))
        if r.status_code != 200 or not zipfile.is_zipfile(io.BytesIO(r.content)):
            continue
        name = "%s %s" % (d.get("process") or "", d.get("name") or d["id"])
        path = os.path.join(folder, re.sub(r"[^\w.-]+", "_", name).strip("_") + ".docx")
        with open(path, "wb") as fh:
            fh.write(r.content)
        out[name.strip()] = path
    return out


def text_of(path, marks=()) -> str:
    """The document's words, paragraphs then table cells, one space apart; each of `marks`
    (version names and ids) and every date replaced, so two versions can be compared."""
    import docx
    document = docx.Document(path)
    words = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            words.extend(cell.text for cell in row.cells)
    text = " ".join(" ".join(words).split())
    for m in sorted({m for m in marks if m}, key=len, reverse=True):
        text = text.replace(m, "<name>")
    return _DATE.sub("<date>", text)


def has_words(ctx, key, vid, words) -> bool:
    """Whether one of version `vid`'s Word files says `words`."""
    wanted = " ".join(str(words).split())
    return any(wanted in text_of(path) for path in fetch(ctx, key, vid).values())


def first_difference(a, b, around=12, most=30):
    """Where two texts first part: (a's words, b's words) around it, the differing words in
    [[ ]]; None when they are the same."""
    wa, wb = a.split(), b.split()
    i = 0
    while i < len(wa) and i < len(wb) and wa[i] == wb[i]:
        i += 1
    if i == len(wa) and i == len(wb):
        return None
    j = 0
    while j < len(wa) - i and j < len(wb) - i and wa[-1 - j] == wb[-1 - j]:
        j += 1
    start = max(0, i - around)

    def show(w):
        mid = w[i:len(w) - j]
        more = " ..." if len(mid) > most else ""
        return ("... " if start else "") + " ".join(w[start:i]) + " [[" + " ".join(mid[:most]) \
            + more + "]]"
    return show(wa), show(wb)


def what_differs(path_a, path_b, name) -> list:
    """doccheck's counted P1 and P2 findings between two documents of one type, worst first:
    ["P2 L4 Lib/Lib/interface: 1 row only one side has ...", ...]."""
    blocks, compare, match, rules, swe3, swe4 = _doccheck()
    profile = swe4 if "SWE.4" in name else swe3
    left = profile.extract(blocks.read(path_a))
    right = profile.extract(blocks.read(path_b))
    result = compare.compare(left, right, profile, match.Aliases(), 5)
    rules.annotate(result, left, right, profile)
    found = sorted((f for f in result.findings if f.counted and f.priority in ("P1", "P2")),
                   key=lambda f: f.sort_key())
    return ["%s %s %s: %s" % (f.priority, f.level, f.path, short(f.summary, 200)) for f in found]


def compare_versions(ctx, key, vid_a, vid_b, *, marks=()) -> dict:
    """The two versions' Word files, compared: {a, b: {name: path}, says: {name: [findings]},
    words: {name: (a's words, b's words)}}. `marks`: more text to take out (project names)."""
    names = set(marks)
    for vid in (vid_a, vid_b):
        _, v = ctx.version(key, vid)
        names |= {vid, v.get("tag")}
    a, b = fetch(ctx, key, vid_a), fetch(ctx, key, vid_b)
    says, words = {}, {}
    for name in sorted(set(a) & set(b)):
        try:
            found = what_differs(a[name], b[name], name)
        except Exception as exc:                    # noqa: BLE001 -- say so, as a difference
            found = ["doccheck could not read them: %s" % short(str(exc), 200)]
        if found:
            says[name] = found
        diff = first_difference(text_of(a[name], names), text_of(b[name], names))
        if diff:
            words[name] = diff
    return {"a": a, "b": b, "says": says, "words": words}


def check_same_documents(ctx, key, vid_a, vid_b, *, what_a, what_b, reqs, words=True,
                         marks=()) -> dict:
    """Three checks: the two versions made the same Word files; they say the same (doccheck: no
    P1 or P2); and, with `words`, word for word the same. The first difference of each document
    is shown, and noted, so a failure says where to look."""
    c = ctx.c
    res = compare_versions(ctx, key, vid_a, vid_b, marks=marks)
    na, nb = set(res["a"]), set(res["b"])
    c.check("%s and %s made the same Word files (%d)" % (what_a, what_b, len(na)),
            bool(na) and na == nb,
            {"only %s" % what_a: sorted(na - nb), "only %s" % what_b: sorted(nb - na)}, reqs=reqs)
    for name, found in res["says"].items():
        c.note("%s -- %s" % (name, "; ".join(found[:3])))
    c.check("they say the same -- doccheck finds no difference at P1 or P2 (%d documents)"
            % len(na & nb), not res["says"],
            {n: f[:4] for n, f in list(res["says"].items())[:4]}, reqs=reqs)
    if words:
        for name, (wa, wb) in res["words"].items():
            c.note("%s, %s: %s" % (name, what_a, short(wa, 400)))
            c.note("%s, %s: %s" % (name, what_b, short(wb, 400)))
        c.check("word for word the same", not res["words"],
                {n: {what_a: short(x, 300), what_b: short(y, 300)}
                 for n, (x, y) in list(res["words"].items())[:3]}, reqs=reqs)
    c.note("the documents: %s" % os.path.join(ctx.log_dir, "documents"))
    return res
