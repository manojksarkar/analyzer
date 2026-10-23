"""The normal scenarios: every requirement of job scheduling once, on short runs without the LLM,
and the queue together with review & update and incremental runs. `--mode normal`.

Each scenario says what it tests (its docstring, printed as it starts), prints its steps, and
names the requirements every check shows (docs/spec/JOB_SCHEDULING_SPEC.md).
"""
from __future__ import annotations

import io
import uuid
import zipfile

from jst_core import (ENDED, HOLDING, WAITING, StopScenario, body_of, ids, is_sentence,
                      scenario, seconds_between, short, when)


def _need_free_slot(ctx, key):
    pv = ctx.preview(key)
    if pv.get("answer") != "start":
        raise StopScenario("the server runs other generations, so a new one would not start now "
                           "(%s): this scenario needs a free slot" % short(pv.get("message")),
                           skip=True)


# ---------------------------------------------------------------------------
@scenario("N01", "Who sees the queues", mode="normal", reqs=("REQ-JS-30",), uses=("A",))
def n01(ctx):
    """A project's two lanes for its members; every project's queues for a super admin only;
    nothing without a sign-in."""
    c = ctx.c
    r = ctx.admin.call("GET", ctx.P("A", "/queues"), auth=False)
    c.check("without a sign-in the queues are refused (401)", r.status_code == 401, r.status_code)
    q = ctx.queues("A")
    lanes = q.get("lanes") or {}
    c.check("a project admin reads its two lanes: generation (limit 3) and update (limit 10)",
            set(lanes) == {"generation", "update"} and lanes["generation"].get("limit") == 3
            and lanes["update"].get("limit") == 10, q)
    c.check("the lanes carry the server's cap (%d, jobs.maxGenerations)" % ctx.cap,
            (q.get("cap") or {}).get("max_generations") == ctx.cap, q.get("cap"),
            reqs=("REQ-JS-05", "REQ-JS-30"))
    if ctx.admin.superuser:
        everything = ctx.all_queues()
        c.check("a super admin reads every project's queues (GET /queues)",
                "projects" in everything and "cap" in everything, short(everything))
    else:
        c.skip("every project's queues", ctx.lacks("superuser"))
    if ctx.dev is None:
        c.skip("a developer's view", ctx.lacks("developer"))
        return
    r = ctx.dev.call("GET", "/queues")
    c.check("a developer is refused every project's queues (403)", r.status_code == 403,
            r.status_code)
    r = ctx.dev.call("GET", ctx.P("A", "/queues"))
    c.check("a developer reads the lanes of a project they are in", r.status_code == 200,
            r.status_code)


# ---------------------------------------------------------------------------
@scenario("N02", "Three answers, and a lane in its order", mode="normal",
          reqs=("REQ-JS-01", "REQ-JS-02", "REQ-JS-06", "REQ-JS-16", "REQ-JS-30", "REQ-JS-31",
                "REQ-JS-35"), uses=("A",))
def n02(ctx):
    """Three generations of one project: the first starts, the others wait in order with their
    place and why; the preview tells what a fourth would get; each starts by itself when the one
    before it ends."""
    c = ctx.c
    _need_free_slot(ctx, "A")
    pv = ctx.preview("A")
    c.check("the preview of an empty lane: 'Starts now'",
            pv.get("answer") == "start" and pv.get("message") == "Starts now", pv,
            reqs=("REQ-JS-31",))
    c.step("submit three generations to A")
    a1, a2 = ctx.submit("A", "order1"), ctx.submit("A", "order2")
    told = ctx.preview("A")
    a3 = ctx.submit("A", "order3")
    for a in (a1, a2, a3):
        if a.code != 202:
            raise StopScenario("a submit was refused: %s" % (a.brief(),))
    c.check("the preview said what the third would get: wait as #2 -- and it did",
            told.get("answer") == "wait" and told.get("position") == 2 == a3.position
            and (told.get("message") or "").startswith("Waits as #2"), (told, a3.brief()),
            reqs=("REQ-JS-31",))
    c.check("the first starts now: running, 'Starts now'",
            a1.status == "running" and a1.message == "Starts now", a1.brief(),
            reqs=("REQ-JS-06",))
    c.check("the second waits as #1, behind the first",
            a2.status == "queued" and a2.position == 1 and a2.message.startswith("Waits as #1")
            and a2.ahead == [a1.job_id], a2.brief(), reqs=("REQ-JS-02", "REQ-JS-06"))
    c.check("the third waits as #2, behind both",
            a3.status == "queued" and a3.position == 2 and a3.message.startswith("Waits as #2")
            and a3.ahead == [a1.job_id, a2.job_id], a3.brief(), reqs=("REQ-JS-06",))
    lane = ctx.lane("A")
    c.check("the lane: the first running, the others waiting in their order",
            ids(lane.get("active")) == [a1.job_id]
            and ids(lane.get("waiting")) == [a2.job_id, a3.job_id], lane,
            reqs=("REQ-JS-02", "REQ-JS-30"))
    waiting = lane.get("waiting") or []
    c.check("each waiting job says where it is in current_activity ('Waiting as #n')",
            [w.get("position") for w in waiting] == [1, 2]
            and all((w.get("current_activity") or "").startswith("Waiting as #%d" % (i + 1))
                    for i, w in enumerate(waiting)),
            [(w.get("position"), w.get("current_activity")) for w in waiting],
            reqs=("REQ-JS-35",))
    j2 = ctx.job("A", a2.job_id)
    c.check("a waiting job has no start time, and carries its lane, priority and place",
            j2.get("started_at") is None and j2.get("lane") == "generation"
            and j2.get("priority") == "normal" and j2.get("position") == 1
            and bool(j2.get("submitted_at")),
            {k: j2.get(k) for k in ("started_at", "lane", "priority", "position",
                                    "submitted_at", "estimated_start")}, reqs=("REQ-JS-30",))
    pv = ctx.preview("A")
    c.check("with three in the lane -- its limit -- the preview says a fourth is refused",
            pv.get("answer") == "refuse" and pv.get("code") == "LANE_FULL", pv,
            reqs=("REQ-JS-31",))
    c.step("follow the three to their ends")
    j1 = ctx.wait_done("A", a1.job_id, what="first")
    c.check("the first completes", j1.get("status") == "complete", j1.get("error_message"))
    j2 = ctx.wait_job("A", a2.job_id, ("running",) + ENDED, what="second", timeout=300)
    gap = seconds_between(j1.get("completed_at"), j2.get("started_at"))
    c.check("the second starts by itself within a minute of the first's end (%s s)"
            % ("?" if gap is None else int(gap)), gap is not None and gap <= 60,
            {"first ended": j1.get("completed_at"), "second started": j2.get("started_at")},
            reqs=("REQ-JS-16",))
    j2 = ctx.wait_done("A", a2.job_id, what="second")
    j3 = ctx.wait_done("A", a3.job_id, what="third")
    starts = [when(j.get("started_at")) for j in (j1, j2, j3)]
    c.check("all three complete, started in their order",
            all(j.get("status") == "complete" for j in (j1, j2, j3)) and None not in starts
            and starts == sorted(starts),
            [(j.get("status"), j.get("started_at")) for j in (j1, j2, j3)])


# ---------------------------------------------------------------------------
@scenario("N03", "The lane's limit; remove and cancel", mode="normal",
          reqs=("REQ-JS-04", "REQ-JS-06", "REQ-JS-16", "REQ-JS-25", "REQ-JS-26", "REQ-JS-31",
                "REQ-JS-35"), uses=("B",))
def n03(ctx):
    """A full lane refuses the next job and names the jobs in it; lowering the limit removes
    nothing; a removed waiting job frees its version name; a cancelled running one stops, its
    unfinished version is deleted, and the next job starts by itself."""
    c = ctx.c
    _need_free_slot(ctx, "B")
    c.step("fill B's generation lane: three jobs, its limit")
    b1, b2, b3 = (ctx.submit("B", "limit%d" % i) for i in (1, 2, 3))
    if [a.code for a in (b1, b2, b3)] != [202, 202, 202]:
        raise StopScenario("a submit was refused: %s" % ([a.brief() for a in (b1, b2, b3)],))
    b4 = ctx.submit("B", "limit4")
    d = b4.detail
    c.check("a fourth is refused: 409 LANE_FULL, naming the running job and the waiting ones",
            b4.code == 409 and b4.error == "LANE_FULL" and ids(d.get("running")) == [b1.job_id]
            and ids(d.get("waiting")) == [b2.job_id, b3.job_id] and d.get("limit") == 3,
            b4.brief(), reqs=("REQ-JS-04", "REQ-JS-06"))
    c.check("its message says what happened and what to do, in words",
            (d.get("message") or "").startswith("Refused: the generation queue of this project "
                                                "is full") and "remove a waiting job" in (
                d.get("message") or ""), d.get("message"), reqs=("REQ-JS-35",))
    pv = ctx.preview("B")
    c.check("the preview says the same: refuse, LANE_FULL",
            pv.get("answer") == "refuse" and pv.get("code") == "LANE_FULL", pv,
            reqs=("REQ-JS-31",))
    r = ctx.set_limits("B", generation_limit=2)
    c.check("an admin lowers the limit to 2", r.code == 200 and r.data.get("generation_limit") == 2,
            r.brief(), reqs=("REQ-JS-04",))
    lane = ctx.lane("B")
    c.check("lowering it removes nothing: the three jobs stay",
            ids(lane.get("active")) == [b1.job_id]
            and ids(lane.get("waiting")) == [b2.job_id, b3.job_id], lane, reqs=("REQ-JS-04",))
    bad = [ctx.set_limits("B", generation_limit=v) for v in (0, 51)]
    c.check("a limit out of 1-50 is 422 INVALID_LIMIT",
            all(a.code == 422 and a.error == "INVALID_LIMIT" for a in bad),
            [a.brief() for a in bad], reqs=("REQ-JS-04", "REQ-JS-35"))
    ctx.set_limits("B", generation_limit=3)
    c.step("remove the last waiting job, and use its version name again")
    r = ctx.act("B", b3.job_id, "cancel")
    c.check("a waiting job cancelled is removed",
            r.code == 200 and r.job.get("status") == "removed", r.brief(), reqs=("REQ-JS-26",))
    again = ctx.submit("B", tag=b3.request["version_tag"])
    c.check("its version name is free again: the same name is accepted",
            again.code == 202, again.brief(), reqs=("REQ-JS-26",))
    if again.job_id:
        ctx.act("B", again.job_id, "cancel")
    c.step("cancel the running job")
    ctx.wait_job("B", b1.job_id, ("running",) + ENDED, phase=1, what="first")
    r = ctx.act("B", b1.job_id, "cancel")
    c.check("cancelling a running job answers 200", r.code == 200, r.brief(),
            reqs=("REQ-JS-25",))
    j1 = ctx.wait_job("B", b1.job_id, ENDED, timeout=120, what="cancelled")
    c.check("it is cancelled, and says by whom",
            j1.get("status") == "cancelled" and (j1.get("error_message") or "").startswith(
                "Cancelled by"), (j1.get("status"), j1.get("error_message")),
            reqs=("REQ-JS-25", "REQ-JS-35"))
    code = ctx.version_gone("B", b1.version_id)
    c.check("its unfinished version is deleted (GET version: 404)", code == 404, code,
            reqs=("REQ-JS-25",))
    j2 = ctx.wait_job("B", b2.job_id, ("running",) + ENDED, timeout=180, what="next")
    c.check("the next waiting job starts by itself",
            j2.get("status") in ("running", "complete"), j2.get("status"), reqs=("REQ-JS-16",))
    ctx.act("B", b2.job_id, "cancel")
    ctx.wait_job("B", b2.job_id, ENDED, timeout=120, what="cancelled")
    r = ctx.act("B", b1.job_id, "cancel")
    c.check("cancelling a job that has ended: 409 JOB_FINISHED, saying how it ended",
            r.code == 409 and r.error == "JOB_FINISHED" and is_sentence(r.detail.get("message")),
            r.brief(), reqs=("REQ-JS-25", "REQ-JS-35"))
    ctx.did.update({"limits", "removed", "cancelled"})


# ---------------------------------------------------------------------------
@scenario("N04", "Checked at submit: names, duplicates, retries, commits, scopes", mode="normal",
          reqs=("REQ-JS-07", "REQ-JS-08", "REQ-JS-09", "REQ-JS-35"), uses=("C",))
def n04(ctx):
    """A version name taken is refused; a generation identical to a waiting one is refused,
    naming it; a request sent again with its Idempotency-Key answers its first job; an unknown
    commit or scope is refused at once -- and none of these takes a version name."""
    c = ctx.c
    _need_free_slot(ctx, "C")
    blocker = ctx.submit("C", "busy")
    waiting = ctx.submit("C", "dup", scope={"type": "group", "names": [ctx.groups[1]]})
    if blocker.code != 202 or waiting.code != 202:
        raise StopScenario("could not set the lane up: %s" % ([blocker.brief(), waiting.brief()],))
    taken = ctx.submit("C", tag=waiting.request["version_tag"],
                       scope={"type": "group", "names": [ctx.groups[2]]})
    c.check("a version name already taken: 409 VERSION_EXISTS",
            taken.code == 409 and taken.error == "VERSION_EXISTS", taken.brief(),
            reqs=("REQ-JS-08",))
    dup = ctx.submit("C", "dup2", scope={"type": "group", "names": [ctx.groups[1]]})
    c.check("the same generation as a waiting one: 409 DUPLICATE_JOB, naming it",
            dup.code == 409 and dup.error == "DUPLICATE_JOB"
            and dup.detail.get("job_id") == waiting.job_id, dup.brief(), reqs=("REQ-JS-08",))
    key = uuid.uuid4().hex
    tag = ctx.tag("retry")
    scope = {"type": "group", "names": [ctx.groups[3 % len(ctx.groups)]]}
    first = ctx.submit("C", tag=tag, scope=scope, idem=key)
    second = ctx.submit("C", tag=tag, scope=scope, idem=key)
    c.check("the same request sent again (one Idempotency-Key): one job",
            first.code == 202 and second.code == 202 and first.job_id == second.job_id,
            [first.brief(), second.brief()], reqs=("REQ-JS-09",))
    unknown = ctx.submit("C", "nocommit", commit="f" * 40)
    c.check("a commit the project does not have: 422 UNKNOWN_COMMIT",
            unknown.code == 422 and unknown.error == "UNKNOWN_COMMIT", unknown.brief(),
            reqs=("REQ-JS-07",))
    names = [v.get("tag") for v in ctx.versions("C")]
    c.check("... and its version name is not taken",
            unknown.request["version_tag"] not in names, names, reqs=("REQ-JS-07", "REQ-JS-08"))
    scope = ctx.submit("C", "noscope", scope={"type": "group", "names": ["Layer9.Nothing"]})
    c.check("a scope the project does not have: 400 INVALID_SCOPE",
            scope.code == 400 and scope.error == "INVALID_SCOPE", scope.brief(),
            reqs=("REQ-JS-07",))
    for a in (taken, dup, unknown, scope):
        c.check("%s says what happened in words" % a.error, is_sentence(a.detail.get("message")),
                a.detail.get("message"), reqs=("REQ-JS-35",))


# ---------------------------------------------------------------------------
@scenario("N05", "Priority in the generation lane", mode="normal",
          reqs=("REQ-JS-11", "REQ-JS-16"), needs=("superuser",), uses=("C",))
def n05(ctx):
    """A High job submitted after two Normal ones goes first, and starts first when the lane
    frees; a super admin changes a waiting job's priority and the order follows it; nobody else
    can, and a running job has none to change."""
    c = ctx.c
    _need_free_slot(ctx, "C")
    # A running job and three waiting: one more than the lane's default limit.
    ctx.set_limits("C", generation_limit=5)
    try:
        _priorities(ctx, c)
    finally:
        ctx.drain("C")
        ctx.set_limits("C", generation_limit=3)
    ctx.did.add("priority")


def _priorities(ctx, c):
    running = ctx.submit("C", "prio-run")
    n1 = ctx.submit("C", "prio-n1")
    n2 = ctx.submit("C", "prio-n2")
    high = ctx.submit("C", "prio-high", priority="high")
    if [a.code for a in (running, n1, n2, high)] != [202] * 4:
        raise StopScenario("could not set the lane up: %s" % (
            [a.brief() for a in (running, n1, n2, high)],))
    c.check("a High job submitted after two Normal ones waits as #1",
            high.status == "queued" and high.position == 1 and high.ahead == [running.job_id],
            high.brief())
    c.check("the lane's order: High first, then the Normal ones oldest first",
            ids(ctx.lane("C").get("waiting")) == [high.job_id, n1.job_id, n2.job_id],
            ids(ctx.lane("C").get("waiting")))
    c.step("cancel the running job: the High one starts next")
    ctx.act("C", running.job_id, "cancel")
    j = ctx.wait_job("C", high.job_id, ("running",) + ENDED, timeout=180, what="high")
    c.check("the High job starts first; the Normal ones still wait",
            j.get("status") in ("running", "complete")
            and ids(ctx.lane("C").get("waiting")) == [n1.job_id, n2.job_id],
            (j.get("status"), ids(ctx.lane("C").get("waiting"))),
            reqs=("REQ-JS-11", "REQ-JS-16"))
    r = ctx.set_priority("C", n2.job_id, "high")
    c.check("a super admin raises the last waiting job to High: it goes first",
            r.code == 200 and r.job.get("priority") == "high"
            and ids(ctx.lane("C").get("waiting")) == [n2.job_id, n1.job_id],
            (r.brief(), ids(ctx.lane("C").get("waiting"))))
    r = ctx.set_priority("C", n2.job_id, "low")
    c.check("... and lowers it to Low: it goes after the Normal one",
            r.code == 200 and ids(ctx.lane("C").get("waiting")) == [n1.job_id, n2.job_id],
            (r.brief(), ids(ctx.lane("C").get("waiting"))))
    r = ctx.set_priority("C", high.job_id, "low")
    c.check("a running job has no priority to change: 409 NOT_WAITING",
            r.code == 409 and r.error == "NOT_WAITING", r.brief())
    r = ctx.set_priority("C", n1.job_id, "urgent")
    c.check("an unknown priority: 422 INVALID_PRIORITY",
            r.code == 422 and r.error == "INVALID_PRIORITY", r.brief())
    if ctx.dev is not None:
        r = ctx.set_priority("C", n1.job_id, "high", api=ctx.dev)
        c.check("a developer sets no priority (403)", r.code == 403, r.brief())
    else:
        c.skip("a developer sets no priority", ctx.lacks("developer"))


# ---------------------------------------------------------------------------
@scenario("N06", "Two lanes at once; the update lane in its order", mode="normal",
          reqs=("REQ-JS-01", "REQ-JS-02", "REQ-JS-08", "REQ-JS-12", "REQ-JS-16", "REQ-JS-22"),
          uses=("A",))
def n06(ctx):
    """A project runs a generation and a Word-file update at the same time; a second update
    waits its turn; the same update again is 'already queued' or joins the running one; an update
    is never paused; updates start in the order they came."""
    c = ctx.c
    v1, v2 = ctx.need_versions("A", 2)[:2]
    _need_free_slot(ctx, "A")
    gen = ctx.submit("A", "with-updates")
    u1 = ctx.reexport("A", v1, {"scope": "all"})
    c.check("an update starts while a generation of the project runs: two lanes",
            gen.code == 202 and gen.status == "running" and u1.code == 202
            and u1.status == "running", [gen.brief(), u1.brief()],
            reqs=("REQ-JS-01", "REQ-JS-02"))
    u2 = ctx.reexport("A", v2, {"scope": "all"})
    c.check("a second update waits as #1 in the update lane",
            u2.code == 202 and u2.status == "queued" and u2.position == 1, u2.brief(),
            reqs=("REQ-JS-01", "REQ-JS-12"))
    again = ctx.reexport("A", v2, {"scope": "all"})
    c.check("the same update again is not added: 200, already queued, its job",
            again.code == 200 and again.data.get("already_queued") is True
            and again.job_id == u2.job_id, again.brief(), reqs=("REQ-JS-08",))
    if ctx.job("A", u1.job_id).get("status") == "running":
        joined = ctx.reexport("A", v1, {"scope": "all"})
        c.check("an update the running one covers joins it: 200, joined",
                joined.code == 200 and joined.data.get("joined") is True
                and joined.job_id == u1.job_id, joined.brief(), reqs=("REQ-JS-08",))
    r = ctx.act("A", u1.job_id, "pause")
    c.check("an update is not paused: 409 NOT_PAUSABLE",
            r.code == 409 and r.error == "NOT_PAUSABLE", r.brief(), reqs=("REQ-JS-22",))
    c.step("follow the updates")
    j1 = ctx.wait_done("A", u1.job_id, what="update 1")
    j2 = ctx.wait_done("A", u2.job_id, what="update 2")
    c.check("both updates complete, the first first",
            j1.get("status") == "complete" and j2.get("status") == "complete"
            and when(j1.get("started_at")) <= when(j2.get("started_at")),
            [(j.get("status"), j.get("started_at")) for j in (j1, j2)],
            reqs=("REQ-JS-12", "REQ-JS-16"))
    ctx.wait_done("A", gen.job_id, what="generation")


# ---------------------------------------------------------------------------
@scenario("N07", "One writer per version", mode="normal",
          reqs=("REQ-JS-03", "REQ-JS-12", "REQ-JS-01", "REQ-JS-35"), uses=("A",))
def n07(ctx):
    """While an export writes components into version v1, an update of v1 waits ('Waiting for
    v1') and keeps its place; an update of another version passes it; v1's update starts when
    the export ends."""
    c = ctx.c
    v1, v2 = ctx.need_versions("A", 2)[:2]
    _need_free_slot(ctx, "A")
    todo = [x["component"] for x in ctx.components("A", v1)
            if x.get("in_model") and x.get("state") == "not_requested"][:2]
    if not todo:
        raise StopScenario("version %s has every component it can have: nothing to export"
                           % v1, skip=True)
    exp = ctx.export("A", v1, todo)
    if exp.code != 202:
        raise StopScenario("the export could not start: %s" % (exp.brief(),))
    c.check("an export of %s into v1 starts, in the generation lane" % ", ".join(todo),
            exp.status == "running", exp.brief(), reqs=("REQ-JS-01",))
    upd = ctx.reexport("A", v1, {"scope": "all"})
    job = ctx.job("A", upd.job_id) if upd.job_id else {}
    c.check("an update of v1 waits for it: 'Waiting for' its version",
            upd.code == 202 and upd.status == "queued"
            and (job.get("current_activity") or "").startswith("Waiting for"),
            (upd.brief(), job.get("current_activity"), job.get("activity_detail")),
            reqs=("REQ-JS-03", "REQ-JS-35"))
    other = ctx.reexport("A", v2, {"scope": "all"})
    c.check("an update of v2 passes it and starts",
            other.code == 202 and other.status == "running", other.brief(),
            reqs=("REQ-JS-12",))
    c.check("v1's update keeps its place: still waiting, #1",
            ctx.job("A", upd.job_id).get("status") == "queued"
            and ids(ctx.lane("A", "update").get("waiting")) == [upd.job_id],
            ctx.lane("A", "update"), reqs=("REQ-JS-12",))
    je = ctx.wait_done("A", exp.job_id, what="export")
    ju = ctx.wait_done("A", upd.job_id, what="v1 update")
    ctx.wait_done("A", other.job_id, what="v2 update")
    c.check("the export completes, then v1's update starts after it and completes",
            je.get("status") == "complete" and ju.get("status") == "complete"
            and when(ju.get("started_at")) >= when(je.get("completed_at")),
            {"export ended": je.get("completed_at"), "update started": ju.get("started_at")},
            reqs=("REQ-JS-03",))
    made = {x["component"]: x.get("state") for x in ctx.components("A", v1)}
    c.check("the exported components are generated in v1",
            all(made.get(x) == "generated" for x in todo), {x: made.get(x) for x in todo})


# ---------------------------------------------------------------------------
@scenario("N08", "The server's cap, and the longest wait first", mode="normal",
          reqs=("REQ-JS-05", "REQ-JS-13", "REQ-JS-16", "REQ-JS-35"), needs=("superuser",),
          uses=("A", "B", "C", "D", "E", "F"))
def n08(ctx):
    """With the cap's generations running, a generation of another project waits for a server
    slot and says so; when a slot frees, the generation that has waited longest starts -- not a
    later one of the project whose job ended. A waiting job's baseline cannot be deleted."""
    c = ctx.c
    cap, keys = ctx.cap, list(ctx.projects)
    if len(keys) < cap + 1:
        raise StopScenario("the cap is %d: this needs %d projects (projects.count)" % (
            cap, cap + 1), skip=True)
    busy = (ctx.all_queues().get("cap") or {}).get("running_generations", 0)
    if busy:
        raise StopScenario("%d generation(s) of other projects run on this server" % busy,
                           skip=True)
    base = ctx.need_versions(keys[cap], 1)[0]
    runners, waiter = keys[:cap], keys[cap]
    c.step("one generation in each of %s: the cap, %d" % (", ".join(runners), cap))
    started = [ctx.submit(k, "cap", scope="long") for k in runners]
    c.check("the first %d start" % cap, all(a.status == "running" for a in started),
            [a.brief() for a in started], reqs=("REQ-JS-05",))
    w = ctx.submit(waiter, "cap-wait", mode="auto", reference=base)
    job = ctx.job(waiter, w.job_id) if w.job_id else {}
    c.check("one more, in another project, waits for a free server slot -- and says so",
            w.code == 202 and w.status == "queued"
            and job.get("current_activity") == "Waiting for a free server slot"
            and "%d of %d generations" % (cap, cap) in (job.get("activity_detail") or ""),
            (w.brief(), job.get("current_activity"), job.get("activity_detail")),
            reqs=("REQ-JS-05", "REQ-JS-35"))
    later = ctx.submit(runners[0], "cap-later")
    c.check("a later job of the first project waits in its lane",
            later.code == 202 and later.status == "queued", later.brief())
    counted = ctx.all_queues().get("cap") or {}
    c.check("GET /queues counts %d of %d generations running" % (cap, cap),
            counted.get("running_generations") == cap and counted.get("max_generations") == cap,
            counted, reqs=("REQ-JS-05", "REQ-JS-30"))
    r = ctx.admin.ask("DELETE", ctx.V(waiter, base))
    c.check("the version a waiting job is built on cannot be deleted: 409 VERSION_NEEDED",
            r.code == 409 and r.error == "VERSION_NEEDED", r.brief())
    ctx._seen(r)
    c.step("cancel the first project's job: its slot goes to the job that waited longest")
    ctx.act(runners[0], started[0].job_id, "cancel")
    jw = ctx.wait_job(waiter, w.job_id, ("running",) + ENDED, timeout=180, what="waiter")
    c.check("the job that waited longest -- of another project -- starts",
            jw.get("status") in ("running", "complete"), jw.get("status"),
            reqs=("REQ-JS-13", "REQ-JS-16"))
    jl = ctx.job(runners[0], later.job_id)
    c.check("the first project's later job still waits, now for a server slot",
            jl.get("status") == "queued", (jl.get("status"), jl.get("current_activity")),
            reqs=("REQ-JS-13",))


# ---------------------------------------------------------------------------
@scenario("N09", "Pause and resume a generation", mode="normal",
          reqs=("REQ-JS-22", "REQ-JS-23", "REQ-JS-24", "REQ-JS-32", "REQ-JS-35"), uses=("A", "B"))
def n09(ctx):
    """A running generation asked to pause stops at its next safe point; paused, it keeps its
    project's lane (the next job of the project waits) but not its server slot (another
    project's job starts); a waiting job is not paused; resumed, it finishes its version."""
    c = ctx.c
    _need_free_slot(ctx, "A")
    p = ctx.submit("A", "pause", scope="long")
    if p.code != 202:
        raise StopScenario("could not start: %s" % (p.brief(),))
    j = ctx.wait_job("A", p.job_id, ("running",) + ENDED, phase=2, what="generation")
    if j.get("status") != "running":
        c.warn("the run ended (%s) before it could be paused" % j.get("status"),
               reqs=("REQ-JS-22",))
        return
    r = ctx.act("A", p.job_id, "pause")
    c.check("asked to pause: 200, pause_requested, 'Pausing at the next safe point…'",
            r.code == 200 and r.job.get("pause_requested") is True
            and (r.job.get("current_activity") or "").startswith("Pausing"), r.brief(),
            reqs=("REQ-JS-22", "REQ-JS-35"))
    j = ctx.wait_job("A", p.job_id, ("paused",) + ENDED, timeout=600, what="pausing")
    if j.get("status") != "paused":
        c.warn("the run reached its end (%s) before a safe point -- a longer run shows it "
               "(complex mode pauses one with the LLM)" % j.get("status"), reqs=("REQ-JS-22",))
        return
    ctx.did.add("paused")
    c.check("it stops at a safe point: paused, saying until when",
            bool(j.get("paused_at")) and (j.get("current_activity") or "").startswith("Paused")
            and "unless resumed" in (j.get("current_activity") or ""),
            (j.get("paused_at"), j.get("current_activity")), reqs=("REQ-JS-22", "REQ-JS-35"))
    nxt = ctx.submit("A", "behind-paused")
    c.check("it keeps the project's lane: the project's next generation waits",
            nxt.code == 202 and nxt.status == "queued" and nxt.ahead == [p.job_id],
            nxt.brief(), reqs=("REQ-JS-24",))
    other = ctx.submit("B", "beside-paused")
    c.check("it does not hold a server slot: another project's generation starts",
            other.code == 202 and other.status == "running", other.brief(),
            reqs=("REQ-JS-24",))
    r = ctx.act("A", nxt.job_id, "pause")
    c.check("a waiting job is not paused: 409 NOT_PAUSABLE",
            r.code == 409 and r.error == "NOT_PAUSABLE", r.brief(), reqs=("REQ-JS-22",))
    r = ctx.act("A", p.job_id, "pause")
    c.check("pausing it again: 409 NOT_PAUSABLE (it is paused)",
            r.code == 409 and r.error == "NOT_PAUSABLE", r.brief(), reqs=("REQ-JS-22",))
    c.step("resume it")
    r = ctx.act("A", p.job_id, "resume")
    c.check("asked to resume: 200, resume_requested or running",
            r.code == 200 and (r.job.get("resume_requested") is True
                               or r.job.get("status") == "running"), r.brief(),
            reqs=("REQ-JS-23",))
    j = ctx.wait_done("A", p.job_id, what="resumed")
    c.check("resumed, it completes its version",
            j.get("status") == "complete", (j.get("status"), j.get("error_message")),
            reqs=("REQ-JS-23",))
    docs = ctx.documents("A", j.get("version_id"))
    c.check("the version has its documents (%d)" % len(docs), bool(docs), len(docs),
            reqs=("REQ-JS-23",))
    r = ctx.act("A", p.job_id, "resume")
    c.check("resuming a job that is not paused: 409 NOT_PAUSED",
            r.code == 409 and r.error == "NOT_PAUSED", r.brief(), reqs=("REQ-JS-23",))
    jn = ctx.wait_job("A", nxt.job_id, ("running",) + ENDED, timeout=300, what="next")
    c.check("the project's next generation starts once the resumed one ends",
            jn.get("status") in ("running", "complete"), jn.get("status"),
            reqs=("REQ-JS-24", "REQ-JS-16"))


# ---------------------------------------------------------------------------
@scenario("N10", "A failed job frees its lane and its name", mode="normal",
          reqs=("REQ-JS-21", "REQ-JS-16", "REQ-JS-35"), uses=("B",))
def n10(ctx):
    """A generation of a commit where a component's folder is gone stops before the parse: the
    job fails with a message, its version (no work in it) is deleted and its name freed, and the
    next job of the project starts by itself."""
    c = ctx.c
    _need_free_slot(ctx, "B")
    rel = _component_dir(ctx, ctx.groups[-1])
    if rel is None:
        raise StopScenario("the last group has no component folder to remove", skip=True)
    branch = "jst-without-%s" % ctx.stamp
    sha = ctx.sample.branch_without(branch, rel)
    c.note("commit %s on branch %s: %s removed" % (sha[:10], branch, rel))
    bad = ctx.submit("B", "fails", commit=sha,
                     scope={"type": "group", "names": [ctx.groups[-1]]})
    nxt = ctx.submit("B", "after-failure")
    if bad.code != 202 or nxt.code != 202:
        raise StopScenario("could not set up: %s" % ([bad.brief(), nxt.brief()],))
    j = ctx.wait_job("B", bad.job_id, ENDED, what="doomed")
    text = j.get("error_message") or ""
    c.check("it fails, with a message", j.get("status") == "failed" and is_sentence(text),
            (j.get("status"), short(text, 400)), reqs=("REQ-JS-21", "REQ-JS-35"))
    code = ctx.version_gone("B", bad.version_id)
    c.check("its version -- nothing stored in it -- is deleted", code == 404, code,
            reqs=("REQ-JS-21",))
    names = [v.get("tag") for v in ctx.versions("B")]
    c.check("... and its name is free", bad.request["version_tag"] not in names, names,
            reqs=("REQ-JS-21",))
    jn = ctx.wait_job("B", nxt.job_id, ("running",) + ENDED, timeout=180, what="next")
    c.check("the next job starts by itself", jn.get("status") in ("running", "complete"),
            jn.get("status"), reqs=("REQ-JS-16", "REQ-JS-21"))
    ctx.did.add("failed")


def _component_dir(ctx, group_id):
    """The folder of the first component of `group_id`, relative to the repository."""
    layer_name, group_name = group_id.split(".", 1)
    for layer in (ctx.cfg.get("projects") or {}).get("architecture_layers") or []:
        if layer.get("name") != layer_name:
            continue
        for g in layer.get("groups") or []:
            if g.get("name") == group_name and g.get("components"):
                files = g["components"][0].get("files") or []
                return files[0] if files else None
    return None


# ---------------------------------------------------------------------------
@scenario("N11", "The command line in the lanes", mode="normal",
          reqs=("REQ-JS-27", "REQ-JS-28", "REQ-JS-29", "REQ-JS-35"), needs=("cli",), uses=("C",))
def n11(ctx):
    """`analyzer.py generate` on this machine records itself as the project's running job; a web
    submission waits behind it and starts when it ends; the web app cannot cancel it; a second
    command-line run is refused, naming the job in the way, and queues nothing; with
    --ignore-queue it runs anyway."""
    c = ctx.c
    _need_free_slot(ctx, "C")
    c.step("start a command-line generate of C")
    run = ctx.cli_generate("C", "cli")

    def cli_job():
        for j in ctx.lane("C").get("active") or []:
            if j.get("origin") == "cli":
                return j
        return None
    found = ctx.wait_for("the command line's job in C's lane", cli_job, 120)
    if found is None:
        out = run.output()
        run.kill()
        raise StopScenario("the command-line run never showed in the lane -- it uses another "
                           "database than the API? Its output: %s" % short(out, 600))
    c.check("the command-line run is the project's running job (origin cli, its runner)",
            found.get("status") == "running" and (found.get("runner") or "").startswith("cli "),
            found, reqs=("REQ-JS-29",))
    web = ctx.submit("C", "behind-cli")
    c.check("a web submission waits behind it",
            web.code == 202 and web.status == "queued" and web.ahead == [found.get("job_id")],
            web.brief(), reqs=("REQ-JS-29",))
    r = ctx.act("C", found["job_id"], "cancel")
    c.check("the web app does not cancel it: 409 CLI_RUN, saying where to stop it",
            r.code == 409 and r.error == "CLI_RUN" and "command line" in (
                r.detail.get("message") or ""), r.brief(), reqs=("REQ-JS-29", "REQ-JS-35"))
    c.step("a second command-line run of C")
    refused = ctx.cli_generate("C", "cli-refused")
    rc = refused.wait(300)
    out = refused.output()
    c.check("it is refused: exit 2, 'Refused: the generation lane ... is taken', naming the job",
            rc == 2 and "Refused: the generation lane of project %s is taken" % ctx.pid("C") in out
            and found["job_id"] in out, (rc, short(out[-800:], 800)),
            reqs=("REQ-JS-28", "REQ-JS-35"))
    c.check("the refused run queued nothing: the lane still waits only the web job",
            ids(ctx.lane("C").get("waiting")) == [web.job_id], ctx.lane("C"),
            reqs=("REQ-JS-27",))
    forced = ctx.cli_generate("C", "cli-forced", extra=("--ignore-queue",))
    both = ctx.wait_for("two running jobs in C", lambda: len(
        ctx.lane("C").get("active") or []) >= 2 or None, 120)
    c.check("--ignore-queue starts it anyway: two jobs run in the lane", both is not None,
            ctx.lane("C"), reqs=("REQ-JS-28",))
    c.step("follow both command-line runs to their end")
    rcs = [run.wait(ctx.timeout), forced.wait(ctx.timeout)]
    c.check("both command-line runs end well (exit 0)", rcs == [0, 0],
            (rcs, short(run.output()[-400:]), short(forced.output()[-400:])))
    jc = ctx.job("C", found["job_id"])
    c.check("its job ends with it: complete", jc.get("status") == "complete", jc.get("status"),
            reqs=("REQ-JS-29",))
    jw = ctx.wait_job("C", web.job_id, ("running",) + ENDED, timeout=300, what="web")
    c.check("the web job behind it starts when it ends", jw.get("status") in (
        "running", "complete"), jw.get("status"), reqs=("REQ-JS-29", "REQ-JS-16"))


# ---------------------------------------------------------------------------
@scenario("N12", "Review & update and incremental runs, through the queue", mode="normal",
          reqs=("REQ-JS-10", "REQ-JS-01", "REQ-JS-02"), uses=("A",))
def n12(ctx):
    """A reviewer's correction makes the Word file out of date; an update -- the update lane --
    brings it in. A version generated from that one carries the correction; and a new commit
    submitted while it runs takes it as its baseline when it starts."""
    c = ctx.c
    v1 = ctx.need_versions("A", 1)[0]
    _need_free_slot(ctx, "A")
    slots = body_of(ctx.admin.call("GET", ctx.V("A", v1, "/slots"), params={
        "slot_kind": "description", "limit": 1000})).get("slots") or []
    target = next((s for s in slots if s.get("shownIn") and not s.get("isOverridden")), None)
    if target is None:
        raise StopScenario("version %s has no description a document prints" % v1, skip=True)
    text = "JST %s: a reviewer's words for %s." % (ctx.stamp, target.get("label") or "it")
    r = ctx.admin.ask("PUT", ctx.V("A", v1, "/overrides/slot"), json={
        "slot_kind": "description", "slot_key": target["slotKey"], "text": text})
    c.check("a description is corrected (review & update)", r.code == 200, r.brief())
    r9 = body_of(ctx.admin.call("GET", ctx.V("A", v1, "/export-readiness")))
    c.check("the Word file is now out of date", r9.get("stale") is True, r9)
    upd = ctx.reexport("A", v1, {"scope": "out_of_date"})
    c.check("an update of the out-of-date Word files is a job of the update lane",
            upd.code == 202 and ctx.job("A", upd.job_id).get("lane") == "update", upd.brief(),
            reqs=("REQ-JS-01",))
    ju = ctx.wait_done("A", upd.job_id, what="update")
    r9 = body_of(ctx.admin.call("GET", ctx.V("A", v1, "/export-readiness")))
    c.check("the update completes and the Word file carries the correction",
            ju.get("status") == "complete" and r9.get("stale") is False,
            (ju.get("status"), r9))
    if ctx.lacks("docx") is None:
        found = _word_has(ctx, "A", v1, text)
        c.check("the corrected words are in a Word file", found, text)
    c.step("generate from v1, and submit a new commit while that runs")
    sha2 = ctx.sample.add_function("next-%s" % ctx.stamp)
    base = ctx.submit("A", "from-v1", mode="auto", reference=v1, scope="first")
    nxt = ctx.submit("A", "new-commit", mode="auto", commit=sha2, scope="first")
    c.check("the new commit's generation waits behind it",
            base.code == 202 and nxt.code == 202 and nxt.status == "queued",
            [base.brief(), nxt.brief()])
    jb = ctx.wait_done("A", base.job_id, what="from v1")
    jn = ctx.wait_done("A", nxt.job_id, what="new commit")
    if jb.get("status") != "complete" or jn.get("status") != "complete":
        raise StopScenario("a generation did not complete: %s" % (
            [jb.get("error_message"), jn.get("error_message")],))
    _, vb = ctx.version("A", base.version_id)
    _, vn = ctx.version("A", nxt.version_id)
    c.check("generated from v1: incremental, its baseline v1",
            vb.get("decision") == "incremental" and vb.get("baseline_version_id") == v1,
            {k: vb.get(k) for k in ("decision", "baseline_version_id")})
    c.check("the new commit's baseline is the version made just before it started",
            vn.get("decision") == "incremental"
            and vn.get("baseline_version_id") == base.version_id,
            {k: vn.get(k) for k in ("decision", "baseline_version_id")}, reqs=("REQ-JS-10",))
    carried = {o.get("slotKey"): o for o in body_of(ctx.admin.call(
        "GET", ctx.V("A", base.version_id, "/overrides"), params={"limit": 1000})).get(
        "overrides") or []}
    got = carried.get(target["slotKey"]) or {}
    c.check("the correction is carried into the version generated from v1",
            got.get("text") == text and got.get("isOverridden") is True, got)


def _word_has(ctx, key, vid, text) -> bool:
    import docx
    for d in ctx.documents(key, vid):
        r = ctx.admin.call("GET", ctx.P(key, "/documents/%s/download" % d["id"]))
        if r.status_code != 200 or not zipfile.is_zipfile(io.BytesIO(r.content)):
            continue
        document = docx.Document(io.BytesIO(r.content))
        words = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                words.extend(cell.text for cell in row.cells)
        if " ".join(text.split()) in " ".join(" ".join(words).split()):
            return True
    return False


# ---------------------------------------------------------------------------
@scenario("N13", "Who may do what: a developer", mode="normal",
          reqs=("REQ-JS-26",), needs=("developer",), uses=("C",))
def n13(ctx):
    """A developer of the project submits a Word-file update for a document they have open, and
    cancels it; they cannot submit a generation, pause, cancel or prioritise someone else's job,
    change a limit, or start the held jobs."""
    c, dev = ctx.c, ctx.dev
    v = ctx.need_versions("C", 1)[0]
    _need_free_slot(ctx, "C")
    r = ctx.submit("C", "dev-gen", api=dev)
    c.check("a developer submits no generation (403)", r.code == 403, r.brief())
    r = ctx.set_limits("C", api=dev, generation_limit=2)
    c.check("a developer changes no limit (403)", r.code == 403, r.brief())
    r = dev.ask("POST", "/queues/held", json={"remove": []})
    c.check("a developer starts no held jobs (403)", r.code == 403, r.brief())
    gen = ctx.submit("C", "admin-gen")
    waiting = ctx.submit("C", "admin-wait")
    r = ctx.act("C", gen.job_id, "pause", api=dev)
    c.check("a developer pauses nothing (403)", r.code == 403, r.brief())
    r = ctx.act("C", gen.job_id, "cancel", api=dev)
    c.check("a developer cancels no one else's job (403)", r.code == 403, r.brief())
    r = ctx.set_priority("C", waiting.job_id, "high", api=dev)
    c.check("a developer sets no priority (403)", r.code == 403, r.brief())
    c.step("an update of the developer's own: correct a description, then update its document")
    slots = body_of(ctx.admin.call("GET", ctx.V("C", v, "/slots"), params={
        "slot_kind": "description", "limit": 1000})).get("slots") or []
    target = next((s for s in slots if s.get("shownIn") and s.get("component")), None)
    if target is None:
        raise StopScenario("no description a document prints in %s" % v, skip=True)
    ctx.admin.call("PUT", ctx.V("C", v, "/overrides/slot"), json={
        "slot_kind": "description", "slot_key": target["slotKey"],
        "text": "JST %s: for the developer's update." % ctx.stamp})
    norm = lambda s: (s or "").strip().replace(" ", "-").casefold()      # noqa: E731
    doc = next((d for d in ctx.documents("C", v)
                if norm(d.get("group")) == norm(target["component"])), None)
    if doc is None:
        raise StopScenario("no document of component %s" % target["component"])
    # The developer's first: an admin's rebuild of every component, running, would cover it.
    own = ctx.reexport("C", v, {"scope": "out_of_date", "document_id": doc["id"]}, api=dev)
    c.check("a developer submits an update for the document they have open (202)",
            own.code == 202 and bool(own.job_id), own.brief(), reqs=("REQ-JS-26",))
    admins = ctx.reexport("C", v, {"scope": "all"})
    if admins.job_id and admins.job_id != own.job_id:
        r = ctx.act("C", admins.job_id, "cancel", api=dev)
        c.check("a developer cancels no one else's update (403)", r.code == 403, r.brief(),
                reqs=("REQ-JS-26",))
    if own.job_id and own.job_id != admins.job_id:
        r = ctx.act("C", own.job_id, "cancel", api=dev)
        c.check("a developer cancels their own update",
                r.code == 200 and r.job.get("status") in ("removed", "cancelled"), r.brief(),
                reqs=("REQ-JS-26",))


# ---------------------------------------------------------------------------
@scenario("N14", "Recorded, told and measured", mode="normal",
          reqs=("REQ-JS-32", "REQ-JS-33", "REQ-JS-34", "REQ-JS-35"))
def n14(ctx):
    """Every step of the jobs above is in the record, with who did it and when; the submitter is
    told when a job starts, ends, pauses or is cancelled; a finished job shows how long it waited
    and ran, phase by phase; every refusal and every ending said what happened in words."""
    c = ctx.c
    events = [e for key in ctx.projects for e in ctx.events(key)]
    me = ctx.admin.uid
    by_kind = {}
    for e in events:
        by_kind.setdefault(e.get("kind"), []).append(e)
    people = {"submitted": True, "cancelled": "cancelled" in ctx.did,
              "removed": "removed" in ctx.did, "limits_changed": "limits" in ctx.did,
              "priority_changed": "priority" in ctx.did, "pause_requested": "paused" in ctx.did,
              "resume_requested": "paused" in ctx.did}
    for kind, expected in people.items():
        if not expected:
            c.skip("the record of %s" % kind, "no scenario of this run did it",
                   reqs=("REQ-JS-33",))
            continue
        mine = [e for e in by_kind.get(kind) or [] if e.get("actor_id") == me]
        c.check("the record has '%s', with who did it and when" % kind,
                bool(mine) and all(e.get("at") for e in mine), len(by_kind.get(kind) or []),
                reqs=("REQ-JS-33",))
    for kind in ("started", "completed") + (("paused", "resumed") if "paused" in ctx.did else ()):
        c.check("the record has '%s' (the system did it)" % kind, bool(by_kind.get(kind)),
                sorted(by_kind), reqs=("REQ-JS-33",))
    pids = {p["id"] for p in ctx.projects.values()}
    notices = [n for n in ctx.notices() if n.get("project_id") in pids]
    types = {n.get("type") for n in notices}
    for kind, did in (("job_started", True), ("job_completed", True),
                      ("job_cancelled", "cancelled" in ctx.did), ("job_paused", "paused" in ctx.did),
                      ("job_failed", "failed" in ctx.did)):
        if not did:
            c.skip("a %s notice" % kind, "no scenario of this run did it", reqs=("REQ-JS-32",))
            continue
        said = [n for n in notices if n.get("type") == kind]
        c.check("the submitter is told: %s, in words" % kind,
                bool(said) and all(is_sentence(n.get("message")) for n in said),
                sorted(types), reqs=("REQ-JS-32", "REQ-JS-35"))
    finished = None
    for job_id, key in ctx.jobs.items():
        j = ctx.job(key, job_id)
        if j.get("status") == "complete" and j.get("lane") == "generation":
            finished = j
            break
    if finished is None:
        c.skip("a finished job's times", "no generation completed", reqs=("REQ-JS-34",))
    else:
        waited = seconds_between(finished.get("submitted_at"), finished.get("started_at"))
        ran = seconds_between(finished.get("started_at"), finished.get("completed_at"))
        phases = finished.get("phases") or []
        c.check("a finished job shows when it was submitted, started and ended "
                "(waited %s s, ran %s s)" % (waited and int(waited), ran and int(ran)),
                waited is not None and waited >= 0 and ran is not None and ran > 0,
                {k: finished.get(k) for k in ("submitted_at", "started_at", "completed_at")},
                reqs=("REQ-JS-34",))
        c.check("... and how long each phase took", bool(phases) and any(
            p.get("duration_seconds") is not None for p in phases), phases, reqs=("REQ-JS-34",))
        c.check("... and its flags (long waits, stuck runs)", isinstance(finished.get("flags"),
                                                                         list),
                finished.get("flags"), reqs=("REQ-JS-34",))
    bad = [x for x in ctx.refusals if not x.get("code") or not is_sentence(x.get("message"))]
    c.check("each of the %d refusals seen carried its code and a message in words"
            % len(ctx.refusals), not bad, bad[:5], reqs=("REQ-JS-35",))
    # The queue's own endings are a headline and its reasons; a failed run's adds the run's last
    # lines for whoever looks into it -- its headline is what the web app shows first.
    queue = [x for x in ctx.endings if x["status"] in ("cancelled", "interrupted")]
    wrong = [x for x in queue if not is_sentence(x["text"].split("\n", 1)[0])
             or any(line and not line.startswith("- ") for line in x["text"].split("\n")[1:])]
    c.check("each of the %d cancels and interruptions seen is a headline, then '- ' lines"
            % len(queue), not wrong, wrong[:5], reqs=("REQ-JS-35",))
    failed = [x for x in ctx.endings if x["status"] == "failed"]
    wrong = [x for x in failed if not is_sentence(x["text"].split("\n", 1)[0])]
    c.check("each of the %d failures seen starts with a headline in words" % len(failed),
            not wrong, wrong[:5], reqs=("REQ-JS-35",))
