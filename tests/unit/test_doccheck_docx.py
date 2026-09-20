"""The whole chain, on a real document: read it, change one thing, compare.

The ladder tests build their trees by hand, which is the right way to test the
comparison and the wrong way to test the reading. These start from a document the
pipeline actually produced, change exactly one cell, row or heading inside the
Word file, and check that the finding comes back out the other end -- extractor,
matcher, policy and report included.

The mutations are deliberately small. A test that rewrites half the document
proves only that the tool noticed something.
"""
import copy
import os
import shutil
import sys

import pytest

_ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
if os.path.join(_ROOT, "tools") not in sys.path:
    sys.path.insert(0, os.path.join(_ROOT, "tools"))

docx = pytest.importorskip("docx", reason="python-docx is needed to edit a .docx")
from docx.oxml.ns import qn                                   # noqa: E402

from doccheck import blocks, compare as comparing, rules, swe3   # noqa: E402

pytestmark = pytest.mark.unit


def _source_document():
    """A generated SWE.3 document with several units and rows, or skip."""
    import glob
    best, best_rows = None, 0
    patterns = [
        os.path.join(_ROOT, "output_review", "all-groups", "*", "software_detailed_design_*.docx"),
        os.path.join(_ROOT, "output", "*", "software_detailed_design_*.docx"),
        os.path.join(_ROOT, "workspaces", "*", "versions", "*", "output", "*",
                     "software_detailed_design_*.docx"),
    ]
    for pattern in patterns:
        for path in sorted(glob.glob(pattern)):
            doc = swe3.extract(blocks.read(path))
            rows = sum(len(u.of_kind("interface"))
                       for c in doc.of_kind("component") for u in c.of_kind("unit"))
            if rows > best_rows:
                best, best_rows = path, rows
    return best


SOURCE = _source_document()
if not SOURCE:
    pytest.skip("no generated SWE.3 document on disk", allow_module_level=True)


# --- editing a Word file ----------------------------------------------------

def _set_text(cell_or_para, value):
    """Replace a cell's or paragraph's text, keeping the first run's formatting."""
    para = cell_or_para.paragraphs[0] if hasattr(cell_or_para, "paragraphs") else cell_or_para
    if hasattr(cell_or_para, "paragraphs"):
        for extra in cell_or_para.paragraphs[1:]:
            extra._p.getparent().remove(extra._p)
    runs = para.runs
    if not runs:
        para.add_run(value)
        return
    runs[0].text = value
    for run in runs[1:]:
        run.text = ""


def _interface_tables(document):
    for table in document.tables:
        header = [c.text.strip().casefold() for c in table.rows[0].cells]
        if any("interface id" in h for h in header):
            yield table, header


def _find_row(document, interface_name):
    """(table, row, header) for the interface table row naming `interface_name`."""
    for table, header in _interface_tables(document):
        name_col = next((i for i, h in enumerate(header) if h.startswith("interface name")), 1)
        for row in table.rows[1:]:
            if row.cells[name_col].text.strip() == interface_name:
                return table, row, header
    raise AssertionError("no interface row named %r" % interface_name)


def _column(header, *starts):
    for i, h in enumerate(header):
        if any(h.startswith(s) for s in starts):
            return i
    raise AssertionError("no column starting with %s in %s" % (starts, header))


def _heading_level(para):
    try:
        name = para.style.name or ""
    except Exception:
        return 0
    return int(name.split()[1]) if name.startswith("Heading ") and name.split()[1].isdigit() else 0


def _delete_unit_section(document, unit_name):
    """Remove a unit's heading and everything under it, up to the next unit."""
    body = document.element.body
    children = list(body.iterchildren())
    start = None
    for i, el in enumerate(children):
        if el.tag != qn("w:p"):
            continue
        from docx.text.paragraph import Paragraph
        para = Paragraph(el, document)
        level = _heading_level(para)
        if start is None:
            if level == 3 and para.text.strip().endswith(unit_name):
                start = i
        elif level and level <= 3:
            for victim in children[start:i]:
                body.remove(victim)
            return
    assert start is not None, "no unit section named %r" % unit_name
    for victim in children[start:]:
        body.remove(victim)


# --- the fixture ------------------------------------------------------------

@pytest.fixture
def mutable(tmp_path):
    """A writable copy of a real document, plus a load-and-compare helper."""
    target = tmp_path / "compared.docx"
    shutil.copyfile(SOURCE, target)

    class Harness:
        path = str(target)
        reference = SOURCE

        def open(self):
            return docx.Document(self.path)

        def save(self, document):
            document.save(self.path)

        def compare(self):
            left = swe3.extract(blocks.read(self.reference))
            right = swe3.extract(blocks.read(self.path))
            result = comparing.compare(left, right, swe3)
            rules.annotate(result, left, right)
            return result

    return Harness()


def _first_function_row():
    """(unit name, interface name) of a function row in the source document."""
    doc = swe3.extract(blocks.read(SOURCE))
    for component in doc.of_kind("component"):
        for unit in component.of_kind("unit"):
            for iface in unit.of_kind("interface"):
                if (iface.fields.get("interfaceType") or "").casefold().startswith("func"):
                    return unit.name, iface.name
    pytest.skip("no function interface row in the source document")


# --- the tests --------------------------------------------------------------

def test_an_untouched_copy_compares_clean(mutable):
    assert mutable.compare().findings == []


def test_a_flipped_direction_cell_reaches_the_report(mutable):
    _unit, name = _first_function_row()
    document = mutable.open()
    table, row, header = _find_row(document, name)
    col = _column(header, "direction")
    before = row.cells[col].text.strip()
    _set_text(row.cells[col], "In" if before != "In" else "Out")
    mutable.save(document)

    result = mutable.compare()
    differs = [f for f in result.findings if f.field == "direction"]
    assert len(differs) == 1, [f.summary for f in result.findings]
    assert differs[0].severity == "high"
    assert name in differs[0].path


def test_rewriting_a_direction_to_an_equivalent_spelling_is_silent(mutable):
    _unit, name = _first_function_row()
    document = mutable.open()
    table, row, header = _find_row(document, name)
    col = _column(header, "direction")
    before = row.cells[col].text.strip()
    _set_text(row.cells[col], before.upper())          # "Out" -> "OUT"
    mutable.save(document)
    assert mutable.compare().findings == []


def test_a_deleted_interface_row_is_reported_as_missing(mutable):
    _unit, name = _first_function_row()
    document = mutable.open()
    table, row, _header = _find_row(document, name)
    row._tr.getparent().remove(row._tr)
    mutable.save(document)

    result = mutable.compare()
    missing = [f for f in result.findings
               if f.kind == "missing" and f.path.endswith(name)]
    assert len(missing) == 1
    counts = [f for f in result.findings if f.kind == "count"]
    assert counts and counts[0].left == counts[0].right + 1


def test_a_deleted_unit_section_reports_the_unit_and_not_its_rows(mutable):
    doc = swe3.extract(blocks.read(SOURCE))
    victim = None
    for component in doc.of_kind("component"):
        for unit in component.of_kind("unit"):
            if len(unit.of_kind("interface")) >= 2:
                victim = unit.name
                break
        if victim:
            break
    if not victim:
        pytest.skip("no unit with two or more interface rows")

    document = mutable.open()
    _delete_unit_section(document, victim)
    mutable.save(document)

    result = mutable.compare()
    below = [f for f in result.findings if "/ %s /" % victim in f.path]
    assert below == [], "the deleted unit leaked %d child findings" % len(below)
    missing = [f for f in result.findings if f.kind == "missing" and f.path.endswith(victim)]
    assert len(missing) == 1 and missing[0].severity == "high"


def test_a_changed_source_destination_is_read_as_a_set(mutable):
    doc = swe3.extract(blocks.read(SOURCE))
    target = None
    for component in doc.of_kind("component"):
        for unit in component.of_kind("unit"):
            for iface in unit.of_kind("interface"):
                if len(iface.fields.get("sourceDest") or []) >= 2:
                    target = iface
                    break
    if target is None:
        pytest.skip("no row with two or more callers")

    document = mutable.open()
    table, row, header = _find_row(document, target.name)
    col = _column(header, "source")
    reversed_text = ", ".join(reversed([s.strip() for s in row.cells[col].text.split(",")]))
    _set_text(row.cells[col], reversed_text)
    mutable.save(document)
    assert mutable.compare().findings == [], "reordering the callers was read as a change"


def test_dropping_a_caller_is_a_difference(mutable):
    doc = swe3.extract(blocks.read(SOURCE))
    target = None
    for component in doc.of_kind("component"):
        for unit in component.of_kind("unit"):
            for iface in unit.of_kind("interface"):
                if len(iface.fields.get("sourceDest") or []) >= 2:
                    target = iface
                    break
    if target is None:
        pytest.skip("no row with two or more callers")

    document = mutable.open()
    table, row, header = _find_row(document, target.name)
    col = _column(header, "source")
    kept = [s.strip() for s in row.cells[col].text.split(",")][:-1]
    _set_text(row.cells[col], ", ".join(kept))
    mutable.save(document)

    differs = [f for f in mutable.compare().findings if f.field == "sourceDest"]
    assert len(differs) == 1 and differs[0].severity == "high"


def test_a_renamed_unit_heading_is_read_as_a_rename(mutable):
    doc = swe3.extract(blocks.read(SOURCE))
    unit_name = doc.of_kind("component")[0].of_kind("unit")[0].name
    document = mutable.open()
    for para in document.paragraphs:
        if _heading_level(para) == 3 and para.text.strip().endswith(unit_name):
            number = para.text.strip()[: -len(unit_name)]
            _set_text(para, "%s%ss" % (number, unit_name))     # Foo -> Foos
            break
    else:
        pytest.skip("could not find the unit heading")
    mutable.save(document)

    result = mutable.compare()
    renamed = [f for f in result.findings if f.kind == "renamed"]
    assert len(renamed) == 1
    assert renamed[0].left == unit_name and renamed[0].right == unit_name + "s"


def test_the_json_report_is_serialisable_and_carries_the_inventory(mutable, tmp_path):
    from doccheck import report
    _unit, name = _first_function_row()
    document = mutable.open()
    table, row, header = _find_row(document, name)
    _set_text(row.cells[_column(header, "direction")], "In/Out")
    mutable.save(document)

    result = mutable.compare()
    import json
    payload = json.loads(report.to_json(result, mutable.reference, mutable.path))
    assert payload["summary"]["findings"] == len(result.findings)
    assert payload["inventory"]["unit"]["reference"] > 0
    assert all("severity" in f for f in payload["findings"])


def test_the_markdown_report_renders_without_the_findings_list_being_empty(mutable):
    from doccheck import report
    _unit, name = _first_function_row()
    document = mutable.open()
    table, row, header = _find_row(document, name)
    _set_text(row.cells[_column(header, "direction")], "In/Out")
    mutable.save(document)

    text = report.markdown(mutable.compare(), "a.docx", "b.docx")
    assert "# Document comparison" in text
    assert "## Ladder" in text
    assert "Direction" in text
